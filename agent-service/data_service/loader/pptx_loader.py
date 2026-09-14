"""Text extraction for text-based PowerPoint presentations."""
from __future__ import annotations

from pathlib import Path
import re
import logging

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from data_service.models.schemas import DocumentChunk

logger = logging.getLogger(__name__)


def _is_slide_number(text: str) -> bool:
    """Do not pollute source text with a standalone page-number placeholder."""
    return bool(re.fullmatch(r"\s*\d{1,3}\s*", text))


def _table_text(shape: object) -> str:
    return "\n".join(
        " | ".join(str(cell.text or "").strip() for cell in row.cells).strip()
        for row in shape.table.rows
    ).strip()


def _native_table_payload(shape: object, *, slide_number: int, table_index: int) -> dict[str, object]:
    """Capture the actual python-pptx Table grid without inferring any values."""
    raw_rows = [
        [str(cell.text or "").strip() for cell in row.cells]
        for row in shape.table.rows
    ]
    columns = raw_rows[0] if raw_rows else []
    rows = raw_rows[1:] if len(raw_rows) > 1 else []
    logger.info(
        "[PPT_SHAPE_DEBUG] slide=%s shape_type=TABLE table_index=%s rows=%s columns=%s",
        slide_number,
        table_index,
        len(rows),
        len(columns),
    )
    return {
        "table_index": table_index,
        "slide": slide_number,
        "columns": columns,
        "rows": rows,
        "text": _table_text(shape),
    }


def _position(shape: object) -> tuple[int, int]:
    return (int(getattr(shape, "top", 0)), int(getattr(shape, "left", 0)))


def _collect_text_from_shapes(
    shapes: object,
    *,
    parent_position: tuple[int, int] | None = None,
) -> list[tuple[tuple[int, int, int, int], str]]:
    """Recursively collect text frames from top-level and grouped shapes.

    A GROUP has no text frame of its own, while its child AUTO_SHAPEs often
    contain the entire visible article body. Child coordinates are local to
    the group, so the group position is retained as the primary read-order key.
    """
    output: list[tuple[tuple[int, int, int, int], str]] = []
    for shape in shapes:
        shape_position = _position(shape)
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            output.extend(_collect_text_from_shapes(shape.shapes, parent_position=shape_position))
            continue
        if not getattr(shape, "has_text_frame", False):
            continue
        text = str(shape.text or "").strip()
        if not text or _is_slide_number(text):
            continue
        group_top, group_left = parent_position or shape_position
        local_top, local_left = shape_position if parent_position is not None else (0, 0)
        output.append(((group_top, group_left, local_top, local_left), text))
    return output


def load_pptx_text(path: str | Path) -> list[DocumentChunk]:
    """Return one provenance-aware chunk for each PPTX slide containing text."""
    pptx_path = Path(path)
    presentation = Presentation(str(pptx_path))
    chunks: list[DocumentChunk] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None:
            title = str(slide.shapes.title.text or "").strip()
        title_shapes: list[tuple[tuple[int, int, int, int], str]] = []
        text_shapes: list[tuple[tuple[int, int, int, int], str]] = []
        table_shapes: list[tuple[tuple[int, int], str]] = []
        native_tables: list[dict[str, object]] = []
        table_index = 0
        chart_count = 0
        for shape in slide.shapes:
            # The title is also a text frame. Collect it once, then recursively
            # inspect ordinary and grouped text shapes without excluding a slide
            # simply because it also contains a table.
            if shape is slide.shapes.title and getattr(shape, "has_text_frame", False):
                text = str(shape.text or "").strip()
                if text and not _is_slide_number(text):
                    top, left = _position(shape)
                    title_shapes.append(((top, left, 0, 0), text))
            elif getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
                text_shapes.extend(_collect_text_from_shapes(shape.shapes, parent_position=_position(shape)))
            elif getattr(shape, "has_text_frame", False):
                text_shapes.extend(_collect_text_from_shapes([shape]))
            if getattr(shape, "has_table", False):
                table_index += 1
                table_payload = _native_table_payload(
                    shape,
                    slide_number=slide_number,
                    table_index=table_index,
                )
                native_tables.append(table_payload)
                text = str(table_payload["text"])
                if text:
                    table_shapes.append((_position(shape), text))
            if getattr(shape, "has_chart", False):
                # Chart data is parsed by pptx_chart_parser. Keep the slide in
                # the unified chunk list even if it has no ordinary text box.
                chart_count += 1
        ordered_parts = title_shapes + text_shapes + [((top, left, 0, 0), text) for (top, left), text in table_shapes]
        parts: list[str] = []
        for _, text in sorted(ordered_parts, key=lambda item: item[0]):
            if text not in parts:
                parts.append(text)
        content = "\n".join(parts).strip()
        if not content and not chart_count:
            continue
        chunks.append(
            DocumentChunk(
                content=content,
                source={"file": pptx_path.name, "slide": slide_number},
                slide=slide_number,
                # Keep each native table's exact extracted text separately.
                # SemanticChunker uses this only for source_text provenance;
                # it never fabricates a table source from parsed/LLM output.
                metadata={
                    "file_type": "pptx",
                    "extractor": "python-pptx",
                    "title": title,
                    "table_texts": [text for _, text in table_shapes],
                    "native_tables": native_tables,
                    "native_chart_count": chart_count,
                },
            )
        )
    return chunks



"""Native chart extraction for text-based PowerPoint presentations."""
from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from pptx import Presentation

from data_service.models.schemas import ChartKnowledge, ChartPoint, ChartSeries, SourceRef

logger = logging.getLogger(__name__)


def _text(value: Any) -> str:
    return " ".join(str(value or "").split())


def _chart_type(chart: Any) -> str:
    """Map the supported Office chart families to stable Chinese labels."""
    raw = str(getattr(chart, "chart_type", ""))
    enum_name = str(getattr(getattr(chart, "chart_type", None), "name", raw)).upper()
    if "BAR" in enum_name or "COLUMN" in enum_name:
        return "柱状图"
    if "LINE" in enum_name:
        return "折线图"
    if "PIE" in enum_name or "DOUGHNUT" in enum_name:
        return "饼图"
    if "BUBBLE" in enum_name:
        return "气泡图"
    return raw or "未知图表"


def _title(chart: Any) -> str:
    try:
        if chart.has_title:
            return _text(chart.chart_title.text_frame.text)
    except Exception:
        logger.debug("Unable to read native chart title", exc_info=True)
    return ""


def _categories(chart: Any) -> list[str]:
    try:
        plots = list(chart.plots)
        if not plots:
            return []
        return [_text(getattr(category, "label", category)) for category in plots[0].categories]
    except Exception:
        logger.debug("Unable to read native chart categories", exc_info=True)
        return []


def _value(value: Any) -> float | str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = _text(value)
    return text or None


def _native_values(series: Any, field_name: str) -> list[float | str | None]:
    """Read a native OOXML value sequence such as xVal or bubbleSize."""
    try:
        container = getattr(series._element, field_name, None)
        if container is None:
            return []
        return [_value(container.pt_v(index)) for index in range(container.ptCount_val)]
    except Exception:
        logger.debug("Unable to read native chart %s values", field_name, exc_info=True)
        return []


def _bubble_points(series: Any, y_values: list[float | str | None]) -> list[ChartPoint]:
    """Keep x/y/size together; x labels alone cannot reconstruct a bubble chart."""
    x_values = _native_values(series, "xVal")
    bubble_sizes = _native_values(series, "bubbleSize")
    point_count = max(len(x_values), len(y_values), len(bubble_sizes))
    if not x_values and not bubble_sizes:
        return []
    return [
        ChartPoint(
            x=x_values[index] if index < len(x_values) else None,
            y=y_values[index] if index < len(y_values) else None,
            bubble_size=bubble_sizes[index] if index < len(bubble_sizes) else None,
        )
        for index in range(point_count)
    ]


def _series(chart: Any) -> list[ChartSeries]:
    output: list[ChartSeries] = []
    try:
        for index, series in enumerate(chart.series, start=1):
            name = _text(getattr(series, "name", "")) or f"系列{index}"
            y_values = [_value(value) for value in series.values]
            output.append(ChartSeries(name=name, values=y_values, points=_bubble_points(series, y_values)))
    except Exception:
        logger.debug("Unable to read native chart series", exc_info=True)
    return output


def extract_pptx_charts(path: str | Path) -> list[ChartKnowledge]:
    """Extract supported native charts from every PPTX slide.

    This intentionally reads only PowerPoint's embedded chart data. Images,
    screenshots and unsupported drawing objects are left untouched; no OCR or
    visual inference is performed.
    """
    pptx_path = Path(path)
    presentation = Presentation(str(pptx_path))
    output: list[ChartKnowledge] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        chart_index = 0
        for shape in slide.shapes:
            if not getattr(shape, "has_chart", False):
                continue
            chart_index += 1
            chart = shape.chart
            kind = _chart_type(chart)
            title = _title(chart) or f"Slide {slide_number} chart {chart_index}"
            categories = _categories(chart)
            series = _series(chart)
            logger.info(
                "[PPT_CHART] slide=%s chart_index=%s type=%s title=%s categories=%s series=%s",
                slide_number, chart_index, kind, title, len(categories), len(series),
            )
            output.append(
                ChartKnowledge(
                    title=title,
                    chart_type=kind,
                    categories=categories,
                    series=series,
                    source=SourceRef(file=pptx_path.name, slide=slide_number),
                    metadata={"chart_index": chart_index, "shape_name": str(getattr(shape, "name", ""))},
                )
            )
    return output



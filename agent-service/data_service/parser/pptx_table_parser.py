"""Rule-based table extraction for text-based PPTX slides."""
from __future__ import annotations

from pathlib import Path
import re

from pptx import Presentation

from data_service.models.schemas import SourceRef, TableData


def _cell_text(cell: object) -> str:
    return " ".join(str(cell.text or "").split())


def _is_stable_table(table: object) -> bool:
    """A table is usable when its declared rectangular grid can be read fully.

    Intentional blank cells are common in grouped business tables, so their
    ratio is not a structural failure signal.
    """
    try:
        rows = list(table.rows)
        column_count = len(table.columns)
        if len(rows) < 2 or column_count < 1:
            return False
        return all(len(list(row.cells)) == column_count for row in rows)
    except Exception:
        return False


_GROUP_HEADER_PATTERN = re.compile(
    r"^(?:前\s*[五十]\s*家.*|TOP\s*(?:5|10).*|第[一二三四五六七八九十0-9]+梯队.*)$",
    re.IGNORECASE,
)


def _clean(value: object) -> str:
    return " ".join(str(value or "").replace("\u200b", "").split())


def _is_group_value(value: str) -> bool:
    return bool(_GROUP_HEADER_PATTERN.match(value.strip()))


def _normalize_group_column(columns: list[str], rows: list[dict[str, str]]) -> tuple[list[str], list[dict[str, str]], dict[str, object]]:
    """Turn a merged first-column ranking label into a fill-down hierarchy field."""
    if not columns or not _is_group_value(columns[0]):
        return columns, rows, {}

    original_column = columns[0]
    normalized_columns = ["排名分组", *columns[1:]]
    current_group = original_column
    group_values: list[str] = []
    normalized_rows: list[dict[str, str]] = []
    for row in rows:
        candidate = _clean(row.get(original_column, ""))
        if candidate:
            # A new non-empty merged-cell value starts the next rank group.
            current_group = candidate
        if current_group and current_group not in group_values:
            group_values.append(current_group)
        normalized_rows.append({
            "排名分组": current_group,
            **{column: row.get(column, "") for column in columns[1:]},
        })
    return normalized_columns, normalized_rows, {
        "group_columns": ["排名分组"],
        "group_values": group_values,
        "original_columns": columns,
        "original_rows": rows,
    }


def extract_pptx_tables(path: str | Path) -> list[TableData]:
    """Extract only stable rectangular tables; complex tables are warned by validation."""
    pptx_path = Path(path)
    presentation = Presentation(str(pptx_path))
    output: list[TableData] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_title = str(slide.shapes.title.text or "").strip() if slide.shapes.title else ""
        for index, shape in enumerate(slide.shapes, start=1):
            if not shape.has_table:
                continue
            table = shape.table
            if not _is_stable_table(table):
                continue
            raw_rows = [[_cell_text(cell) for cell in row.cells] for row in table.rows]
            rows = [row for row in raw_rows if any(row)]
            if len(rows) < 2 or not rows[0]:
                continue
            columns = [_clean(value) or f"column_{position}" for position, value in enumerate(rows[0], start=1)]
            data_rows = [
                {column: _clean(row[position]) if position < len(row) else "" for position, column in enumerate(columns)}
                for row in rows[1:]
            ]
            columns, data_rows, metadata = _normalize_group_column(columns, data_rows)
            output.append(
                TableData(
                    title=slide_title or f"Slide {slide_number} table {index}",
                    columns=columns,
                    rows=data_rows,
                    source=SourceRef(file=pptx_path.name, slide=slide_number),
                    metadata=metadata,
                )
            )
    return output



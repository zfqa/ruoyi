"""单元格解析."""
from typing import Dict, Any
import openpyxl
from app.models.cell import ExcelCell
from app.excel.formatter import format_display_value, infer_data_type
from app.excel.merged_cell import build_merge_map, get_merge_info


def parse_cells(ws_formula, ws_value, sheet_name: str) -> list:
    """解析 sheet 中所有单元格，返回 ExcelCell 列表."""
    merge_map = build_merge_map(ws_formula)
    cells = []
    max_row = ws_formula.max_row
    max_col = ws_formula.max_column

    # 使用同一个迭代器遍历两份 sheet
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cf = ws_formula.cell(r, c)
            cv = ws_value.cell(r, c)

            formula = None
            cached_value = None
            if cf.data_type == "f":
                formula = cf.value
                cached_value = cv.value
            raw_value = cf.value

            merge_info = get_merge_info(merge_map, r, c)

            cell = ExcelCell(
                sheet_name=sheet_name,
                row=r,
                column=c,
                coordinate=cf.coordinate,
                raw_value=raw_value,
                display_value=format_display_value(
                    cached_value if formula else raw_value,
                    cf.number_format,
                ),
                data_type=infer_data_type(raw_value, formula),
                formula=formula,
                cached_value=cached_value,
                number_format=cf.number_format,
                is_merged=merge_info["is_merged"],
                merged_range=merge_info["merged_range"],
                merged_source_cell=merge_info["merged_source_cell"],
                merged_source_value=merge_info["merged_source_value"],
                row_hidden=ws_formula.row_dimensions[r].hidden or False,
                column_hidden=ws_formula.column_dimensions[openpyxl.utils.get_column_letter(c)].hidden or False,
                hyperlink=cf.hyperlink.target if cf.hyperlink else None,
                comment=cf.comment.text if cf.comment else None,
            )
            cells.append(cell)
    return cells

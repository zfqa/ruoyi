"""合并单元格处理."""
from typing import Dict, Optional, Any
import openpyxl


def build_merge_map(ws: openpyxl.worksheet.worksheet.Worksheet) -> Dict[str, dict]:
    """构建单元格坐标 -> 合并单元格信息的映射."""
    merge_map = {}
    for merged_range in ws.merged_cells.ranges:
        min_row, min_col, max_row, max_col = (
            merged_range.min_row,
            merged_range.min_col,
            merged_range.max_row,
            merged_range.max_col,
        )
        top_left = ws.cell(min_row, min_col)
        source_coord = top_left.coordinate
        source_value = top_left.value
        range_str = merged_range.coord
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                coord = openpyxl.utils.get_column_letter(c) + str(r)
                merge_map[coord] = {
                    "is_source": (r == min_row and c == min_col),
                    "source_coord": source_coord,
                    "source_value": source_value,
                    "range": range_str,
                }
    return merge_map


def get_merge_info(merge_map: Dict[str, dict], row: int, col: int) -> dict:
    coord = openpyxl.utils.get_column_letter(col) + str(row)
    info = merge_map.get(coord)
    if not info:
        return {"is_merged": False, "merged_range": None, "merged_source_cell": None, "merged_source_value": None}
    return {
        "is_merged": True,
        "merged_range": info["range"],
        "merged_source_cell": info["source_coord"],
        "merged_source_value": info["source_value"],
    }

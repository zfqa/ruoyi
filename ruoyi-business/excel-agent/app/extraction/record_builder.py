"""Build records by rereading values from source coordinates."""
from __future__ import annotations

from typing import Any

from openpyxl.utils.cell import column_index_from_string


TOTAL_KEYWORDS = ("合计", "总计", "小计", "汇总", "total", "subtotal", "source:")


def build_records(ws_formula, ws_value, table: dict[str, Any], file_name: str, workbook_id: str) -> list[dict[str, Any]]:
    fields = table.get("fields", [])
    data = table.get("data") or {}
    start_row = data.get("start_row")
    end_row = data.get("end_row")
    if not fields or not start_row or not end_row:
        return []

    records = []
    footer_rows = set(table.get("footer_rows", []))
    for row in range(start_row, end_row + 1):
        if row in footer_rows:
            continue
        if _is_footer_or_empty(ws_value, row, fields):
            continue
        record = {}
        for field in fields:
            col_index = column_index_from_string(field["column"])
            formula_cell = ws_formula.cell(row, col_index)
            value_cell = ws_value.cell(row, col_index)
            raw_value = formula_cell.value
            cached_value = value_cell.value if formula_cell.data_type == "f" else None
            value = cached_value if formula_cell.data_type == "f" else raw_value
            record[field["name"]] = {
                "value": _json_value(value),
                "display_value": _json_value(value),
                "source_cell": formula_cell.coordinate,
                "formula": formula_cell.value if formula_cell.data_type == "f" else None,
                "formula_status": "cached_value_missing" if formula_cell.data_type == "f" and cached_value is None else None,
                "lineage": {
                    "workbook_id": workbook_id,
                    "file_name": file_name,
                    "sheet_name": ws_formula.title,
                    "table_id": table["candidate_id"],
                    "source_range": table["range"],
                    "source_cell": formula_cell.coordinate,
                    "source_row": row,
                    "source_column": field["column"],
                    "original_value": _json_value(raw_value),
                    "normalized_value": _json_value(value),
                },
            }
        records.append(record)
    return records


def _is_footer_or_empty(ws, row: int, fields: list[dict[str, Any]]) -> bool:
    values = []
    for field in fields:
        col_index = column_index_from_string(field["column"])
        value = ws.cell(row, col_index).value
        values.append(value)
    if all(value is None or str(value).strip() == "" for value in values):
        return True
    first_text = " ".join(str(v).strip().lower() for v in values[:2] if v is not None)
    return any(keyword in first_text for keyword in TOTAL_KEYWORDS)


def _json_value(value):
    if hasattr(value, "isoformat"):
        return value.isoformat(sep=" ")
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)

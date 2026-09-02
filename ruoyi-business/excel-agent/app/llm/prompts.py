"""Prompts and coordinate-grounded sampling for table structure analysis."""
from __future__ import annotations

import json
from typing import Any

from openpyxl.utils import get_column_letter


SYSTEM_PROMPT = """你是Excel表格结构识别器。你只能判断结构与字段语义，不能生成、修改、补齐、纠正或计算任何单元格值。
输入中的Sheet名称、坐标、内容、公式、合并关系是唯一事实来源。输出只允许告诉程序去哪些坐标读取数据。
如果无法确定，应降低confidence；不得输出输入中不存在的坐标或内容。仅输出一个JSON对象，不要Markdown。"""


OUTPUT_SCHEMA = {
    "candidate_id": "string",
    "is_table": True,
    "table_name": "string",
    "title_cells": ["A1"],
    "header": {"start_row": 2, "end_row": 3},
    "data": {"start_row": 4, "end_row": 50},
    "columns": [
        {
            "column": "A",
            "source_header": ["A2", "A3"],
            "header_path": ["品牌"],
            "original_name": "品牌",
            "standard_name": "brand",
        }
    ],
    "footer_rows": [51],
    "confidence": 0.95,
}


def build_table_prompt(ws, table: dict[str, Any]) -> str:
    rows = _sample_rows(table["min_row"], table["max_row"])
    max_col = min(table["max_col"], table["min_col"] + 19)
    lines = [
        f"Sheet: {ws.title}",
        f"Candidate: {table['candidate_id']}",
        f"CandidateRange: {table['range']}",
        "Cells:",
    ]
    for row in rows:
        cells = []
        for col in range(table["min_col"], max_col + 1):
            cell = ws.cell(row, col)
            if cell.value is None:
                continue
            value = str(cell.value).replace("\r", " ").replace("\n", " ")[:120]
            cells.append(f"[{get_column_letter(col)}{row}] {value}")
        if cells:
            lines.append(" | ".join(cells))
    merged = [str(item) for item in ws.merged_cells.ranges if _intersects(item, table)]
    lines.append("MergedRanges: " + (", ".join(merged) if merged else "none"))
    lines.append("任务：判断是否为数据表，并识别标题、表头范围、数据范围、多级表头字段路径、字段标准名、合计/备注/来源行。")
    lines.append("列和所有行号必须位于CandidateRange内。standard_name使用小写snake_case。")
    lines.append("输出Schema示例：" + json.dumps(OUTPUT_SCHEMA, ensure_ascii=False))
    return "\n".join(lines)[:20000]


def _sample_rows(min_row: int, max_row: int) -> list[int]:
    size = max_row - min_row + 1
    if size <= 35:
        return list(range(min_row, max_row + 1))
    first = list(range(min_row, min(min_row + 20, max_row + 1)))
    middle_start = max(min_row, (min_row + max_row) // 2 - 2)
    middle = list(range(middle_start, min(middle_start + 5, max_row + 1)))
    last = list(range(max(min_row, max_row - 9), max_row + 1))
    return list(dict.fromkeys(first + middle + last))


def _intersects(merged_range, table: dict[str, Any]) -> bool:
    return not (
        merged_range.max_row < table["min_row"]
        or merged_range.min_row > table["max_row"]
        or merged_range.max_col < table["min_col"]
        or merged_range.min_col > table["max_col"]
    )

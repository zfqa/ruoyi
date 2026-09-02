"""表格候选区域与解析结果模型."""
from dataclasses import dataclass, field
from typing import List, Dict, Any, Optional


@dataclass
class TableRegion:
    candidate_id: str
    sheet: str
    min_row: int
    max_row: int
    min_col: int
    max_col: int
    density: float = 0.0
    text_ratio: float = 0.0
    number_ratio: float = 0.0
    merge_ratio: float = 0.0

    def range_str(self) -> str:
        from openpyxl.utils import get_column_letter
        return f"{get_column_letter(self.min_col)}{self.min_row}:{get_column_letter(self.max_col)}{self.max_row}"

"""Table structure validation."""
from __future__ import annotations

from openpyxl.utils.cell import column_index_from_string


def validate_table(table: dict, ws) -> list[str]:
    errors = []
    header = table.get("header")
    data = table.get("data") or {}
    if not header:
        errors.append("header_missing")
        return errors
    if header["start_row"] < 1 or header["end_row"] > ws.max_row:
        errors.append("header_out_of_range")
    if data.get("start_row", 0) < header["end_row"]:
        errors.append("data_starts_before_header_end")
    if data.get("end_row", 0) > ws.max_row:
        errors.append("data_end_out_of_range")
    for field in table.get("fields", []):
        try:
            col = column_index_from_string(field["column"])
        except ValueError:
            errors.append(f"invalid_column:{field.get('column')}")
            continue
        if col > ws.max_column:
            errors.append(f"column_out_of_range:{field['column']}")
        if not field.get("source_cell"):
            errors.append(f"field_source_missing:{field.get('name')}")
    return errors

import csv
import os
import sys
import tempfile
import unittest

from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from app.pipeline import parse_workbook


class PipelineTest(unittest.TestCase):
    def test_csv_uses_the_same_table_pipeline(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "market.csv")
            with open(path, "w", encoding="utf-8-sig", newline="") as handle:
                writer = csv.writer(handle)
                writer.writerow(["Year", "Panel Maker", "Qty (000)"])
                writer.writerow([2025, "BOE", 10.5])
                writer.writerow([2026, "TCL CSOT", 12])

            result = parse_workbook(path, include_raw_cells=True)

        self.assertEqual("market.csv", result["file_name"])
        self.assertEqual(1, len(result["sheets"]))
        table = result["sheets"][0]["tables"][0]
        self.assertEqual(["year", "panel_maker", "quantity_000"], [field["name"] for field in table["fields"]])
        self.assertEqual(2, len(table["records"]))
        self.assertEqual(10.5, table["records"][0]["quantity_000"]["value"])

    def test_merged_cells_and_missing_formula_cache_are_preserved(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "formula.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Market"
            sheet.merge_cells("A1:C1")
            sheet["A1"] = "Market data"
            sheet.append(["Year", "Panel Maker", "Qty (000)"])
            sheet.append([2025, "BOE", 10])
            sheet.append([2026, "TCL CSOT", "=C3+2"])
            workbook.save(path)

            result = parse_workbook(path, include_raw_cells=True)

        sheet_result = result["sheets"][0]
        self.assertIn("A1:C1", sheet_result["merged_ranges"])
        table = sheet_result["tables"][0]
        formula_value = table["records"][1]["quantity_000"]
        self.assertEqual("=C3+2", formula_value["formula"])
        self.assertEqual("cached_value_missing", formula_value["formula_status"])

    def test_metrics_use_all_rows_before_record_preview_is_truncated(self):
        with tempfile.TemporaryDirectory() as directory:
            path = os.path.join(directory, "shipment.xlsx")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Shipment"
            sheet.append([
                "Year", "Quarter", "Panel Maker", "Original Specification",
                "Product", "Technology", "Qty (000)",
            ])
            sheet.append([2024, "Q1", "Tianma", "Automobile monitor", "Control panel", "LTPS TFT LCD", 100])
            sheet.append([2025, "Q1", "Tianma", "Automobile monitor", "Control panel", "LTPS TFT LCD", 120])
            sheet.append([2025, "Q2", "AUO", "Automobile monitor", "CID", "a-Si TFT LCD", 80])
            workbook.save(path)

            result = parse_workbook(path, max_records_per_table=1)

        table = result["sheets"][0]["tables"][0]
        self.assertEqual(1, len(table["records"]))
        self.assertEqual(200, result["computed_metrics"]["market"]["values"]["2025"])
        self.assertEqual(3, result["computed_metrics"]["scope"]["source_row_count"])


if __name__ == "__main__":
    unittest.main()

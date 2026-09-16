from __future__ import annotations
import io
import hashlib
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.models.schemas import ContextItem
from app.services.context_category import normalize_context_category

CATEGORY_KEYWORDS: dict[str, list[str]] = {
    "macro_policy": [
        "政策", "法规", "条例", "指南", "征求意见", "工信部", "商务部", "财政部", "市场监管",
        "补贴", "税收", "关税", "欧盟", "美国政府", "日本政府", "十五五", "准入", "碳边境", "cbam",
    ],
    "personnel": [
        "任命", "履新", "离职", "辞任", "卸任", "接任", "出任", "升任", "人事", "ceo", "总裁",
        "董事长", "首席", "负责人", "换帅", "高管",
    ],
    "strategy": [
        "战略", "布局", "投资", "工厂", "产能", "出海", "海外", "合作", "合资", "重组", "私有化",
        "收缩", "扩张", "渠道", "品牌", "产品规划", "组织调整", "本地化生产", "并购",
    ],
    "industry_chain": [
        "产业链", "智能驾驶", "自动驾驶", "robotaxi", "电池", "固态电池", "芯片", "座舱", "显示",
        "micro led", "oled", "lcd", "供应链", "tier1", "传感器", "激光雷达", "充电", "零部件",
    ],
    "competition": [
        "竞争", "专利", "诉讼", "收购", "股权", "市场份额", "份额", "竞品", "竞争追踪", "定点",
        "客户", "供应商", "排名", "价格战", "反内卷",
    ],
}


def classify_text(text: str) -> str:
    t = text.lower().strip()
    # 已有周报/报告类资料优先遵循章节标题，避免正文中的政策词把产业链内容误分到政策章节。
    first = t[:120]
    if re.search(r"(?:^|\n)\s*1\.1.*市场.*观察", first):
        return "other"
    if re.search(r"(?:^|\n)\s*1\.2.*宏观政策", first):
        return "macro_policy"
    if re.search(r"(?:^|\n)\s*2\.1.*人事", first):
        return "personnel"
    if re.search(r"(?:^|\n)\s*2\.2.*战略", first):
        return "strategy"
    if re.search(r"(?:^|\n)\s*3(?:\.1)?.*产业链", first):
        return "industry_chain"
    if re.search(r"(?:^|\n)\s*4(?:\.1)?.*竞争追踪", first):
        return "competition"
    scores = {k: sum(2 if kw.lower() in t[:160] else 1 for kw in kws if kw.lower() in t) for k, kws in CATEGORY_KEYWORDS.items()}
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "other"


def _clean(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text or "")).strip()


def _title_from_text(text: str, fallback: str) -> str:
    lines = [x.strip(" •-—\t") for x in text.splitlines() if x.strip()]
    if lines:
        title = lines[0]
        if len(title) <= 90:
            return title
    sentence = re.split(r"[。！？!?\n]", text.strip())[0]
    return (sentence[:88] + "…") if len(sentence) > 90 else (sentence or fallback)


def _split_long_text(text: str, max_chars: int = 1200) -> list[str]:
    text = _clean(text)
    if not text:
        return []
    paragraphs = [p.strip() for p in text.split("\n") if p.strip()]
    chunks: list[str] = []
    current = ""
    for p in paragraphs:
        if len(current) + len(p) + 1 <= max_chars:
            current = (current + "\n" + p).strip()
        else:
            if current:
                chunks.append(current)
            if len(p) <= max_chars:
                current = p
            else:
                for i in range(0, len(p), max_chars):
                    chunks.append(p[i:i + max_chars])
                current = ""
    if current:
        chunks.append(current)
    return chunks


def items_from_text(text: str, source_name: str, locator: str = "", category: str | None = None) -> list[ContextItem]:
    items: list[ContextItem] = []
    section_only = re.compile(r"^(?:\d+(?:\.\d+)?[.、]?\s*)?(?:行业全景|企业行动|人事调整|人事任命|战略调整与布局|产业链观察|竞争追踪|宏观政策动态)$")
    for idx, chunk in enumerate(_split_long_text(text), 1):
        # PPT/Word中经常存在纯章节分隔页；它们不是事实资料，不进入周报事件库。
        lines = [x.strip() for x in chunk.splitlines() if x.strip() and not x.strip().isdigit()]
        normalized = " ".join(lines).strip()
        toc_terms = ["行业全景", "宏观政策", "企业行动", "人事调整", "人事任命", "战略调整与布局", "产业链观察", "竞争追踪"]
        toc_hits = sum(1 for term in toc_terms if term in normalized)
        if not normalized or (len(lines) <= 2 and section_only.match(normalized)) or (toc_hits >= 4 and len(normalized) < 240):
            continue
        loc = locator or (f"片段{idx}" if idx > 1 else "")
        items.append(ContextItem(
            # 手工录入以用户明确选择的分类为准；文件解析未传分类时仍使用自动识别。
            category=normalize_context_category(category or classify_text(chunk)),
            title=_title_from_text(chunk, source_name),
            content=chunk,
            source_name=source_name,
            locator=loc,
            content_sha256=hashlib.sha256(chunk.encode("utf-8")).hexdigest(),
        ))
    return items


def _parse_docx(path: Path) -> list[ContextItem]:
    doc = Document(path)
    blocks: list[str] = []
    current: list[str] = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        is_heading = "heading" in style or "标题" in style
        if is_heading and current:
            blocks.append("\n".join(current))
            current = [text]
        else:
            current.append(text)
    if current:
        blocks.append("\n".join(current))
    items: list[ContextItem] = []
    for idx, block in enumerate(blocks or ["\n".join(p.text for p in doc.paragraphs)], 1):
        items.extend(items_from_text(block, path.name, f"段落组{idx}"))
    return items


def _shape_texts(shape) -> list[str]:
    texts: list[str] = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            texts.append(text)
    # 参考周报大量正文位于组合形状中，需要递归遍历 group.shapes。
    if hasattr(shape, "shapes"):
        for child in shape.shapes:
            texts.extend(_shape_texts(child))
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if line:
                texts.append(line)
    return texts


def _parse_pptx(path: Path) -> list[ContextItem]:
    prs = Presentation(path)
    items: list[ContextItem] = []
    for idx, slide in enumerate(prs.slides, 1):
        texts: list[str] = []
        for shape in slide.shapes:
            texts.extend(_shape_texts(shape))
        joined = "\n".join(texts)
        if joined.strip():
            items.extend(items_from_text(joined, path.name, f"第{idx}页"))
    return items


def _parse_pdf(path: Path) -> list[ContextItem]:
    reader = PdfReader(str(path))
    items: list[ContextItem] = []
    for idx, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            items.extend(items_from_text(text, path.name, f"第{idx}页"))
    return items


def _row_text(values: Iterable[object]) -> str:
    return " | ".join(str(v).strip() for v in values if v is not None and str(v).strip())


def _parse_xlsx(path: Path) -> list[ContextItem]:
    wb = load_workbook(path, data_only=True, read_only=True)
    items: list[ContextItem] = []
    for ws in wb.worksheets:
        rows = []
        for ridx, row in enumerate(ws.iter_rows(values_only=True), 1):
            text = _row_text(row)
            if text:
                rows.append(text)
            if ridx >= 500:
                break
        if rows:
            items.extend(items_from_text("\n".join(rows), path.name, f"工作表:{ws.title}"))
    return items


def _parse_csv(path: Path) -> list[ContextItem]:
    raw = path.read_bytes()
    for enc in ["utf-8-sig", "utf-8", "gb18030", "gbk"]:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    try:
        df = pd.read_csv(io.StringIO(text), sep=None, engine="python").head(500)
        lines = [_row_text(df.columns.tolist())]
        lines += [_row_text(row) for row in df.itertuples(index=False, name=None)]
        text = "\n".join(lines)
    except Exception:
        pass
    return items_from_text(text, path.name, "CSV")


def parse_context_file(path: Path) -> list[ContextItem]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        raw = path.read_bytes()
        for enc in ["utf-8-sig", "utf-8", "gb18030", "gbk"]:
            try:
                return items_from_text(raw.decode(enc), path.name)
            except UnicodeDecodeError:
                continue
        return items_from_text(raw.decode("utf-8", errors="replace"), path.name)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pptx":
        return _parse_pptx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in {".xlsx", ".xlsm"}:
        return _parse_xlsx(path)
    if suffix == ".csv":
        return _parse_csv(path)
    raise ValueError("行业资料仅支持 txt/md/docx/pptx/pdf/xlsx/xlsm/csv")

from __future__ import annotations
import csv
import hashlib
import io
import re
from pathlib import Path
from typing import Any
from collections import Counter

import chardet
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.models.schemas import ParseIssue, SheetMeta
from app.services.field_mapping import (
    clean_header, known_keyword_count, map_headers_with_confidence, field_mapping_confidence, deterministic_field_validation, DIMENSION_FIELDS,
)
from app.services.sheet_semantics import (
    classify_sheet_semantics,
    choose_semantic_dimension_column,
    choose_semantic_value_column,
    is_period_header,
    detect_period_columns,
    normalize_period_header,
    reshape_wide_metric_table,
)
from app.services.value_parser import parse_numeric, extract_yoy, parse_date_like, normalize_text
from app.services.parsers.template_classifier import classify_template
from app.services.parsers.annual_block_parser import parse_annual_month_blocks


CANONICAL_NUMERIC_FIELDS = ["production", "sales", "retail_sales", "wholesale", "domestic_sales", "domestic_wholesale", "export", "inventory"]
YOY_BY_VALUE_FIELD = {
    "production": "yoy_production",
    "sales": "yoy_sales",
    "retail_sales": "yoy_retail_sales",
    "wholesale": "yoy_wholesale",
    "domestic_sales": "yoy_domestic_sales",
    "domestic_wholesale": "yoy_domestic_wholesale",
    "export": "yoy_export",
}
AMBIGUOUS_DIM_HEADERS = {"名称", "对象", "项目", "规格", "类别", "分类", "类型", "明细", "项目名称", "对象名称", "name", "item", "type"}


def file_hash(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def expand_merged_cells(ws: Worksheet) -> dict[tuple[int, int], Any]:
    """读取worksheet物理结构，并将合并单元格根值向下/向右填充。"""
    values: dict[tuple[int, int], Any] = {}
    for row in ws.iter_rows():
        for cell in row:
            values[(cell.row, cell.column)] = cell.value
    for merged in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged.bounds
        root_val = ws.cell(min_row, min_col).value
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                values[(r, c)] = root_val
    return values


def row_values(values: dict[tuple[int, int], Any], row_idx: int, max_col: int) -> list[Any]:
    return [values.get((row_idx, c)) for c in range(1, max_col + 1)]


def longest_contiguous_nonempty(row: list[Any]) -> int:
    best = cur = 0
    for v in row:
        if normalize_text(v):
            cur += 1
            best = max(best, cur)
        else:
            cur = 0
    return best


def detect_header_row(values: dict[tuple[int, int], Any], max_row: int, max_col: int) -> int | None:
    """遍历前100行识别最可信表头。

    甲方固定模板的关键特征是表头包含大量 202401/202402/... 年月列，
    因此周期列数量被作为高权重信号，不再只依赖“销量/产量”等字段词。
    """
    best_row = None
    best_score = -1.0
    scan_to = min(max_row, 100)
    for r in range(1, scan_to + 1):
        vals = row_values(values, r, max_col)
        non_empty = sum(1 for v in vals if normalize_text(v))
        if non_empty < 2:
            continue
        kw = known_keyword_count(vals)
        contig = longest_contiguous_nonempty(vals)
        text_count = sum(1 for v in vals if normalize_text(v) and not isinstance(v, (int, float)))
        period_count = sum(1 for v in vals if is_period_header(v))
        score = kw * 10 + contig * 2 + text_count + period_count * 7
        if period_count >= 6:
            score += 25
        elif period_count >= 2:
            score += 10
        if non_empty <= 2 and kw < 2 and period_count < 2:
            score -= 10
        if score > best_score:
            best_score = score
            best_row = r
    return best_row if best_score >= 6 else None


def infer_sheet_type(headers: list[Any], sample_rows: list[list[Any]], semantic_type: str | None = None) -> str:
    if semantic_type and semantic_type != "unknown_table":
        return semantic_type
    text = "|".join(normalize_text(v) for v in headers)
    text += "|" + "|".join(normalize_text(v) for row in sample_rows[:5] for v in row)
    if any(k in text for k in ["Omdia", "面板", "显示", "尺寸", "技术路线", "出货面积"]):
        return "vehicle_display_tracker"
    if any(k in text for k in ["销量", "产量", "库存", "出口", "批发"]):
        return "vehicle_market"
    return "unknown_table"


def dedupe_headers(headers: list[str]) -> list[str]:
    seen: dict[str, int] = {}
    out = []
    for h in headers:
        base = h or "未命名列"
        seen[base] = seen.get(base, 0) + 1
        out.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return out




def profile_period_cells(df: pd.DataFrame) -> dict[str, Any]:
    """统计固定年月宽表的源数据规模。

    这里特意区分：
    - source_data_rows：Excel表头下面的原始明细行；
    - period_valid_cells：可转成数值的“明细×月份”数据点；
    - standardized rows：宽转长之后真正进入统计引擎的记录数。
    这样前端不会再把 512 行Excel误显示成 10979 行。
    """
    period_cols = detect_period_columns(df.columns)
    if not period_cols:
        return {
            "period_column_count": 0,
            "period_start": None,
            "period_end": None,
            "period_cells_total": 0,
            "period_valid_cells": 0,
            "period_placeholder_cells": 0,
            "period_placeholder_breakdown": {},
        }

    valid = 0
    placeholder_counter: Counter[str] = Counter()
    for col in period_cols:
        for value in df[col].tolist():
            if parse_numeric(value) is not None:
                valid += 1
                continue
            text = normalize_text(value)
            placeholder_counter[text if text else "(空白)"] += 1

    normalized_periods = [normalize_period_header(c) for c in period_cols]
    return {
        "period_column_count": len(period_cols),
        "period_start": normalized_periods[0] if normalized_periods else None,
        "period_end": normalized_periods[-1] if normalized_periods else None,
        "period_cells_total": len(df) * len(period_cols),
        "period_valid_cells": valid,
        "period_placeholder_cells": sum(placeholder_counter.values()),
        "period_placeholder_breakdown": dict(placeholder_counter.most_common(10)),
    }


def _unique_business_values(series: pd.Series) -> set[str]:
    values: set[str] = set()
    for value in series.tolist():
        if value is None or (isinstance(value, float) and pd.isna(value)):
            continue
        text = normalize_text(value)
        if text:
            values.add(text)
    return values


def build_dimension_integrity_profile(
    source_df: pd.DataFrame,
    standardized_df: pd.DataFrame,
    mapping: dict[str, str],
) -> dict[str, Any]:
    """比较源表与标准化后的维度唯一值数量，防止静默丢失。

    该校验只针对明确映射的维度字段。N/A是合法业务分类，不作为空值删除。
    """
    profile: dict[str, Any] = {}
    for src, std in mapping.items():
        if std not in DIMENSION_FIELDS or src not in source_df.columns or std not in standardized_df.columns:
            continue
        source_vals = _unique_business_values(source_df[src])
        std_vals = _unique_business_values(standardized_df[std])
        missing = sorted(source_vals - std_vals)
        profile[std] = {
            "source_column": src,
            "source_unique_count": len(source_vals),
            "standardized_unique_count": len(std_vals),
            "missing_after_standardization": missing[:30],
            "status": "完整" if not missing else "有缺失",
        }
    return profile


def _effective_dimension_semantics(
    semantic_dimension: str | None,
    semantic_dimension_confidence: float,
    mapping: dict[str, str],
    field_confidences: dict[str, float],
) -> tuple[str | None, float, str | None]:
    """固定多维明细表优先以表头实际映射结果描述维度。

    例如“集团/整车厂品牌/车种/车型/动力总成”同时存在时，文件名里的
    “整车厂”不应把整张表误标成单一OEM表。
    """
    mapped_dims = [(src, std) for src, std in mapping.items() if std in DIMENSION_FIELDS]
    distinct_dims = list(dict.fromkeys(std for _, std in mapped_dims))
    if len(distinct_dims) >= 2:
        confs = [field_confidences.get(src, 0.0) for src, _ in mapped_dims]
        confidence = min(confs) if confs else 0.0
        return "multi_dimension", round(confidence, 3), "multi_dimension_detail"
    if len(distinct_dims) == 1 and max((field_confidences.get(src, 0.0) for src, _ in mapped_dims), default=0.0) >= 0.95:
        return distinct_dims[0], round(max(field_confidences.get(src, 0.0) for src, _ in mapped_dims), 3), None
    return semantic_dimension, semantic_dimension_confidence, None


def _coerce_market_dimensions(std: pd.DataFrame) -> pd.DataFrame:
    if "market" in std.columns and "vehicle_type" not in std.columns:
        std["vehicle_type"] = std["market"].where(
            std["market"].astype(str).str.contains("乘用|商用|新能源|燃油", na=False), None
        )
    if "vehicle_type" in std.columns and "market" not in std.columns:
        std["market"] = std["vehicle_type"]
    return std


def normalize_dataframe(
    df: pd.DataFrame,
    mapping: dict[str, str],
    source_file: str,
    sheet_name: str,
    header_row: int | None,
    detected_metric: str | None = None,
    detected_dimension: str | None = None,
    table_type: str | None = None,
    semantic_confidence: float = 0.0,
    source_rows: list[int] | None = None,
    source_column_indexes: dict[str, int] | None = None,
) -> pd.DataFrame:
    rename = {src: std for src, std in mapping.items() if src in df.columns}
    std = df.rename(columns=rename).copy()

    for field in CANONICAL_NUMERIC_FIELDS:
        if field in std.columns:
            yoy_field = YOY_BY_VALUE_FIELD.get(field)
            if yoy_field and yoy_field not in std.columns:
                std[yoy_field] = std[field].map(extract_yoy)
            std[field] = std[field].map(parse_numeric)

    for c in [c for c in std.columns if c.startswith("yoy_")]:
        std[c] = std[c].map(parse_numeric)

    if "time_period" in std.columns:
        std["time_period"] = std["time_period"].map(parse_date_like)

    std = _coerce_market_dimensions(std)
    std["source_file"] = source_file
    std["source_sheet"] = sheet_name
    std["source_metric_type"] = detected_metric
    std["source_dimension"] = detected_dimension
    std["source_table_type"] = table_type
    std["source_semantic_confidence"] = semantic_confidence
    if "source_row" not in std.columns:
        if source_rows is not None and len(source_rows) == len(std):
            std["source_row"] = source_rows
        elif header_row is not None:
            std["source_row"] = [header_row + 1 + i for i in range(len(std))]

    # 长表也记录指标值所在的原始单元格。宽表转换已提供更精确的周期列坐标，
    # 因而只在尚无source_cell时补充，绝不覆盖上游血缘。
    if "source_cell" not in std.columns and "source_row" in std.columns and source_column_indexes:
        metric_source = next(
            (src for src, target in mapping.items() if target == detected_metric and src in source_column_indexes),
            None,
        )
        if metric_source:
            column_index = source_column_indexes[metric_source]
            from openpyxl.utils import get_column_letter
            std["source_column"] = str(metric_source)
            std["source_column_index"] = column_index
            std["source_cell"] = [
                f"{get_column_letter(column_index)}{int(row)}" if pd.notna(row) else None
                for row in std["source_row"]
            ]
    return std


def _apply_semantic_dimension_mapping(
    df: pd.DataFrame,
    mapping: dict[str, str],
    detected_dimension: str | None,
    dimension_confidence: float,
) -> tuple[dict[str, str], list[str]]:
    """当Sheet/文件名明确“品牌/车型/OEM/动力”时，给左侧对象列赋标准维度。"""
    notes: list[str] = []
    if not detected_dimension or dimension_confidence != 1.0:
        return mapping, notes
    if detected_dimension in mapping.values():
        return mapping, notes

    out = dict(mapping)
    # 先允许强语义覆盖“类别/分类/名称”等泛化表头的错误泛匹配。
    for src, std in list(out.items()):
        if clean_header(src) in {clean_header(x) for x in AMBIGUOUS_DIM_HEADERS} and std in {"market", "vehicle_type"}:
            out[src] = detected_dimension
            notes.append(f"根据Sheet/文件维度语义，将泛化列“{src}”解释为{detected_dimension}")
            return out, notes

    col, reason = choose_semantic_dimension_column(df, detected_dimension, set(out.keys()))
    if col:
        out[col] = detected_dimension
        notes.append(f"根据Sheet/文件维度语义，将列“{col}”解释为{detected_dimension}（{reason}）")
    else:
        notes.append(f"已识别维度{detected_dimension}，但{reason}，未自动指定对象列")
    return out, notes


def _semantic_transform(
    df: pd.DataFrame,
    mapping: dict[str, str],
    detected_metric: str | None,
    semantic_confidence: float,
    detected_dimension: str | None,
    dimension_confidence: float,
    source_rows: list[int] | None = None,
    source_column_indexes: dict[str, int] | None = None,
) -> tuple[pd.DataFrame, dict[str, str], str | None, str, list[str]]:
    """根据文件名/Sheet名说明型模板完成维度赋义、宽表展开和数值列赋义。"""
    value_column: str | None = None
    transform_mode = "standard"
    notes: list[str] = []

    mapping, dim_notes = _apply_semantic_dimension_mapping(df, mapping, detected_dimension, dimension_confidence)
    notes.extend(dim_notes)

    if not detected_metric:
        return df, mapping, value_column, transform_mode, notes

    # 固定宽表：对象列 + 202401 + 202402 + ...
    reshaped, reshape_notes = reshape_wide_metric_table(
        df, detected_metric, mapping,
        source_rows=source_rows,
        source_column_indexes=source_column_indexes,
    )
    if reshaped is not None:
        return reshaped, mapping, None, "wide_to_long", notes + reshape_notes

    # 长表：Sheet名表示产量/销量，但数值列表头只有“数值/数量/本期”。
    mapped_metrics = set(mapping.values())
    if detected_metric not in mapped_metrics and semantic_confidence == 1.0:
        mapped_sources = set(mapping.keys())
        col, reason = choose_semantic_value_column(df, mapped_sources)
        if col:
            mapping = dict(mapping)
            mapping[col] = detected_metric
            value_column = col
            transform_mode = "semantic_value_column"
            notes.append(f"根据表类型“{detected_metric}”将列“{col}”解释为该指标（{reason}）")
        else:
            notes.append(f"工作表已识别为{detected_metric}，但{reason}，未自动指定数值列")
    return df, mapping, value_column, transform_mode, notes



NUMERIC_PLACEHOLDERS = {"", "-", "—", "–", "n/a", "na", "null", "none", "缺失"}


def _period_validation(df: pd.DataFrame, mapping: dict[str, str]) -> dict[str, Any]:
    period_cols = detect_period_columns(df.columns)
    if period_cols:
        normalized = [normalize_period_header(c) for c in period_cols]
        ok = all(bool(re.fullmatch(r"20\d{2}-\d{2}|20\d{2}-Q[1-4]|20\d{2}", str(x))) for x in normalized)
        return {"score": 1.0 if ok else 0.0, "mode": "wide", "count": len(period_cols), "normalized": normalized[:36]}
    src = next((src for src, std in mapping.items() if std == "time_period" and src in df.columns), None)
    if src:
        values = [x for x in df[src].tolist() if normalize_text(x)]
        parsed = [parse_date_like(x) for x in values]
        ok = bool(values) and all(x is not None for x in parsed)
        return {"score": 1.0 if ok else 0.0, "mode": "long", "count": len(values), "invalid_count": sum(x is None for x in parsed)}
    # 静态单期长表可没有时间字段；模板本身仍可解析，但不允许宣称周期校验通过。
    return {"score": 0.0, "mode": "none", "count": 0}


def _value_validation(df: pd.DataFrame, mapping: dict[str, str], metric: str | None) -> dict[str, Any]:
    period_cols = detect_period_columns(df.columns)
    bad: list[dict[str, Any]] = []
    checked = 0
    if period_cols:
        for col in period_cols:
            for idx, value in enumerate(df[col].tolist(), start=1):
                text = normalize_text(value).strip()
                if clean_header(text) in NUMERIC_PLACEHOLDERS:
                    continue
                checked += 1
                if parse_numeric(value) is None:
                    bad.append({"column": str(col), "row_offset": idx, "value": text[:80]})
                    if len(bad) >= 20:
                        break
            if len(bad) >= 20:
                break
    else:
        metric_fields = {"production", "sales", "retail_sales", "wholesale", "domestic_sales", "domestic_wholesale", "export", "inventory"}
        sources = [src for src, std in mapping.items() if std in metric_fields and src in df.columns]
        if metric and not sources:
            # 可能稍后由generic value列赋义；调用方会传更新后的mapping。
            sources = [src for src, std in mapping.items() if std == metric and src in df.columns]
        for col in sources:
            for idx, value in enumerate(df[col].tolist(), start=1):
                text = normalize_text(value).strip()
                if clean_header(text) in NUMERIC_PLACEHOLDERS:
                    continue
                checked += 1
                if parse_numeric(value) is None:
                    bad.append({"column": str(col), "row_offset": idx, "value": text[:80]})
                    if len(bad) >= 20:
                        break
    # 无可检查数值时不宣称100%。
    return {"score": 1.0 if checked > 0 and not bad else 0.0, "checked": checked, "invalid_samples": bad}


def _template_validation(df: pd.DataFrame, mapping: dict[str, str], field_confidences: dict[str, float], semantic, effective_dimension: str | None, dimension_integrity: dict[str, Any]) -> dict[str, Any]:
    field_v = deterministic_field_validation(df, mapping, field_confidences)
    period_v = _period_validation(df, mapping)
    value_v = _value_validation(df, mapping, semantic.metric)
    metric_score = 1.0 if semantic.confidence == 1.0 else 0.0
    # 多指标长表由表头精确字段可通过；宽表必须有单一metric。
    direct_metric_count = len({std for std in mapping.values() if std in CANONICAL_NUMERIC_FIELDS})
    if metric_score == 0.0 and semantic.metric_status == "PASS_MULTI" and direct_metric_count >= 2:
        metric_score = 1.0
    dim_ok = all(item.get("status") == "完整" for item in dimension_integrity.values()) if dimension_integrity else True
    dimension_score = 1.0 if dim_ok and (effective_dimension is not None or any(std in DIMENSION_FIELDS for std in mapping.values())) else 0.0
    conflicts = list(semantic.conflicts or [])
    core_scores = [metric_score, field_v["score"], period_v["score"], value_v["score"], dimension_score]
    score = 1.0 if all(x == 1.0 for x in core_scores) and not conflicts else 0.0
    if score == 1.0 and field_v.get("unlabeled_data_columns"):
        status = "CORE_PASS_AUX_REVIEW"
    else:
        status = "PASS" if score == 1.0 else "REVIEW"
    return {
        "score": score,
        "status": status,
        "metric_score": metric_score,
        "field_score": field_v["score"],
        "period_score": period_v["score"],
        "value_score": value_v["score"],
        "dimension_score": dimension_score,
        "field": field_v,
        "period": period_v,
        "value": value_v,
        "conflicts": conflicts,
        "note": "100%仅表示注册核心模板全部确定性通过；无表头辅助列单独列示，不参与核心指标计算。" if status == "CORE_PASS_AUX_REVIEW" else "",
    }

def parse_excel(path: Path, source_name: str | None = None) -> tuple[pd.DataFrame, list[SheetMeta], list[ParseIssue]]:
    source_file = source_name or path.name
    wb = load_workbook(path, data_only=True)
    all_frames: list[pd.DataFrame] = []
    metas: list[SheetMeta] = []
    issues: list[ParseIssue] = []

    for ws in wb.worksheets:
        values = expand_merged_cells(ws)
        merged_ranges = [str(rng) for rng in ws.merged_cells.ranges]
        decision = classify_template(source_file, ws.title, ws)
        if decision.parser_id == "annual_month_blocks_parser" and decision.metric:
            std, meta, block_issues = parse_annual_month_blocks(ws, source_file, decision, merged_ranges)
            if not std.empty:
                std["source_sheet_name"] = ws.title
                all_frames.append(std)
            metas.append(meta)
            issues.extend(block_issues)
            continue
        header_row = decision.header_rows[0] if decision.structure == "yyyymm_wide" and decision.header_rows else detect_header_row(values, ws.max_row, ws.max_column)
        if not header_row:
            metas.append(SheetMeta(source_file=source_file, sheet_name=ws.title, sheet_type="unknown_table", merged_ranges=merged_ranges))
            issues.append(ParseIssue(level="warning", code="HEADER_NOT_FOUND", message="未识别到可信表头，已跳过该工作表", sheet=ws.title))
            continue

        headers_raw = row_values(values, header_row, ws.max_column)
        headers = dedupe_headers([normalize_text(h) or f"列{idx+1}" for idx, h in enumerate(headers_raw)])
        rows = []
        source_rows: list[int] = []
        blank_streak = 0
        for r in range(header_row + 1, ws.max_row + 1):
            vals = row_values(values, r, ws.max_column)
            if sum(1 for v in vals if normalize_text(v)) == 0:
                blank_streak += 1
                if blank_streak >= 3:
                    break
                continue
            blank_streak = 0
            rows.append(vals)
            source_rows.append(r)
        df = pd.DataFrame(rows, columns=headers).dropna(how="all")
        source_column_indexes = {str(header): idx + 1 for idx, header in enumerate(headers)}
        period_profile = profile_period_cells(df)

        unnamed_with_data = []
        for col in df.columns:
            if str(col).startswith("列") and df[col].map(lambda x: bool(normalize_text(x))).any():
                unnamed_with_data.append(str(col))
        if unnamed_with_data:
            issues.append(ParseIssue(
                level="warning", code="UNNAMED_DATA_COLUMNS", sheet=ws.title, row=header_row,
                message="发现表头为空但下方存在数据的列：" + "、".join(unnamed_with_data) + "。这些列保留在标准化明细中，但因缺少业务名称不自动映射，也不计入已命名字段映射置信度。",
            ))

        semantic = classify_sheet_semantics(source_file, ws.title, headers, rows)
        # New router: registered structure/name classification is deterministic and
        # takes precedence over old heuristic confidence values.
        if decision.status == "PASS" and decision.metric:
            semantic.metric = decision.metric
            semantic.confidence = 1.0
            semantic.evidence.extend([x for x in decision.evidence if x not in semantic.evidence])
        mapping, field_confidences = map_headers_with_confidence(headers)
        original_mapping = dict(mapping)

        direct_metrics = {v for v in mapping.values() if v in CANONICAL_NUMERIC_FIELDS}
        if semantic.metric and direct_metrics and semantic.metric not in direct_metrics:
            issues.append(ParseIssue(
                level="warning", code="SEMANTIC_CONFLICT", sheet=ws.title, row=header_row,
                message=f"Sheet/文件语义识别为 {semantic.metric}，但表头直接映射到 {sorted(direct_metrics)}；系统保留明确表头，不强制覆盖，请人工复核。",
            ))

        transformed_df, mapping, value_column, transform_mode, transform_notes = _semantic_transform(
            df, mapping, semantic.metric, semantic.confidence, semantic.dimension, semantic.dimension_confidence,
            source_rows=source_rows,
            source_column_indexes=source_column_indexes,
        )
        # 语义规则给“名称/规格/类型”等泛化列完成明确赋义时，也记录该字段的可审计置信度。
        for src, std_name in mapping.items():
            if src not in field_confidences:
                if semantic.dimension and std_name == semantic.dimension and semantic.dimension_confidence >= 0.95:
                    field_confidences[src] = semantic.dimension_confidence
                elif semantic.metric and std_name == semantic.metric and semantic.confidence >= 0.95:
                    field_confidences[src] = semantic.confidence
        for note in transform_notes:
            issues.append(ParseIssue(level="info", code="SEMANTIC_MAPPING", message=note, sheet=ws.title, row=header_row))

        mapping_conf = field_mapping_confidence(df, mapping, field_confidences)
        effective_dimension, effective_dimension_conf, inferred_table_type = _effective_dimension_semantics(
            semantic.dimension, semantic.dimension_confidence, mapping, field_confidences
        )
        if inferred_table_type:
            semantic.evidence.append(
                f"表头同时明确映射到多个业务维度，按固定多维明细表处理；字段映射置信度{mapping_conf:.1%}"
            )
        if mapping_conf != 1.0:
            issues.append(ParseIssue(
                level="warning", code="FIELD_MAPPING_REVIEW", sheet=ws.title, row=header_row,
                message="固定模板已命名业务字段存在未注册字段；不会用模糊概率冒充100%，请人工确认或补充模板注册表。",
            ))

        sheet_type = inferred_table_type or infer_sheet_type(headers, rows, semantic.table_type)
        if not mapping and transform_mode == "standard":
            issues.append(ParseIssue(level="warning", code="FIELD_MAPPING_LOW", message="未匹配到标准字段，请检查模板表头或文件/Sheet命名", sheet=ws.title, row=header_row))

        std = normalize_dataframe(
            transformed_df, mapping, source_file, ws.title, header_row,
            detected_metric=semantic.metric,
            detected_dimension=effective_dimension,
            table_type=sheet_type,
            semantic_confidence=semantic.confidence,
            source_rows=source_rows if transform_mode != "wide_to_long" else None,
            source_column_indexes=source_column_indexes,
        )
        if not std.empty:
            std["source_metric_scope"] = decision.metric_scope
            std["source_metric_label"] = decision.metric_label or decision.metric
            std["source_parser_id"] = decision.parser_id
            std["source_row_kind"] = "detail"
        dimension_integrity = build_dimension_integrity_profile(df, std, mapping)
        for dim_name, item in dimension_integrity.items():
            if item.get("status") != "完整":
                issues.append(ParseIssue(
                    level="warning", code="DIMENSION_INTEGRITY", sheet=ws.title, row=header_row, column=dim_name,
                    message=f"维度{dim_name}源表唯一值{item.get('source_unique_count')}个，标准化后{item.get('standardized_unique_count')}个；存在静默丢失风险。",
                ))
        validation = _template_validation(df, mapping, field_confidences, semantic, effective_dimension, dimension_integrity)
        if validation["score"] != 1.0:
            issues.append(ParseIssue(
                level="warning", code="TEMPLATE_VALIDATION_REVIEW", sheet=ws.title, row=header_row,
                message=f"固定模板确定性校验未全部通过：指标={validation['metric_score']:.0%}、字段={validation['field_score']:.0%}、周期={validation['period_score']:.0%}、数值={validation['value_score']:.0%}、维度={validation['dimension_score']:.0%}。未通过项不会自动标成100%。",
            ))
        elif validation["status"] == "CORE_PASS_AUX_REVIEW":
            issues.append(ParseIssue(
                level="info", code="CORE_TEMPLATE_PASS_AUX_REVIEW", sheet=ws.title, row=header_row,
                message="核心固定模板确定性校验100%通过；另有无表头辅助数据列，已保留但不参与核心指标计算，请按需人工确认。",
            ))
        # v10 多Sheet隔离：保留来源Sheet，后续分析可按工作表独立切换
        std['source_sheet_name'] = ws.title
        all_frames.append(std)
        metas.append(SheetMeta(
            source_file=source_file,
            sheet_name=ws.title,
            sheet_type=sheet_type,
            detected_metric=semantic.metric,
            detected_dimension=effective_dimension,
            semantic_confidence=semantic.confidence,
            dimension_confidence=effective_dimension_conf,
            field_mapping_confidence=mapping_conf,
            field_confidences=field_confidences,
            metric_validation_score=validation["metric_score"],
            dimension_validation_score=validation["dimension_score"],
            period_validation_score=validation["period_score"],
            value_validation_score=validation["value_score"],
            template_validation_score=validation["score"],
            template_validation_status=validation["status"],
            validation_details=validation,
            unmapped_named_columns=validation["field"].get("unmapped_named_columns", []),
            unlabeled_data_columns=validation["field"].get("unlabeled_data_columns", []),
            dimension_integrity=dimension_integrity,
            semantic_evidence=semantic.evidence,
            value_column=value_column,
            transform_mode=transform_mode,
            header_row=header_row,
            original_headers=headers,
            field_mapping=mapping or original_mapping,
            rows_detected=len(std),
            physical_rows=ws.max_row,
            physical_columns=ws.max_column,
            source_data_rows=len(df),
            period_column_count=period_profile["period_column_count"],
            period_start=period_profile["period_start"],
            period_end=period_profile["period_end"],
            period_cells_total=period_profile["period_cells_total"],
            period_valid_cells=period_profile["period_valid_cells"],
            period_placeholder_cells=period_profile["period_placeholder_cells"],
            period_placeholder_breakdown=period_profile["period_placeholder_breakdown"],
            merged_ranges=merged_ranges,
            parser_id=decision.parser_id,
            metric_scope=decision.metric_scope,
            metric_label=decision.metric_label or decision.metric,
            structure_type=decision.structure,
            block_count=len(decision.header_rows),
        ))

    if not all_frames:
        return pd.DataFrame(), metas, issues
    combined = pd.concat(all_frames, ignore_index=True, sort=False)
    return combined, metas, issues


def detect_csv_dialect(raw: bytes) -> tuple[str, csv.Dialect]:
    enc_guess = chardet.detect(raw[:100000]).get("encoding") or "utf-8"
    for enc in [enc_guess, "utf-8-sig", "gb18030", "gbk", "utf-8"]:
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=[",", "\t", ";", "|"])
    except Exception:
        dialect = csv.get_dialect("excel")
    return text, dialect


def parse_csv(path: Path, source_name: str | None = None) -> tuple[pd.DataFrame, list[SheetMeta], list[ParseIssue]]:
    source_file = source_name or path.name
    raw = path.read_bytes()
    text, dialect = detect_csv_dialect(raw)
    sep = dialect.delimiter
    df = pd.read_csv(io.StringIO(text), sep=sep, keep_default_na=False, na_values=[""])
    df.columns = dedupe_headers([normalize_text(c) for c in df.columns])

    rows = df.head(8).values.tolist()
    period_profile = profile_period_cells(df)
    semantic = classify_sheet_semantics(source_file, "CSV", list(df.columns), rows)
    mapping, field_confidences = map_headers_with_confidence(df.columns)
    transformed_df, mapping, value_column, transform_mode, transform_notes = _semantic_transform(
        df, mapping, semantic.metric, semantic.confidence, semantic.dimension, semantic.dimension_confidence
    )
    for src, std_name in mapping.items():
        if src not in field_confidences:
            if semantic.dimension and std_name == semantic.dimension and semantic.dimension_confidence >= 0.95:
                field_confidences[src] = semantic.dimension_confidence
            elif semantic.metric and std_name == semantic.metric and semantic.confidence >= 0.95:
                field_confidences[src] = semantic.confidence
    mapping_conf = field_mapping_confidence(df, mapping, field_confidences)
    effective_dimension, effective_dimension_conf, inferred_table_type = _effective_dimension_semantics(
        semantic.dimension, semantic.dimension_confidence, mapping, field_confidences
    )

    issues: list[ParseIssue] = []
    if mapping_conf != 1.0:
        issues.append(ParseIssue(level="warning", code="FIELD_MAPPING_REVIEW", message="CSV存在未注册命名字段，确定性字段映射未通过；不会用模糊概率冒充100%。"))
    for note in transform_notes:
        issues.append(ParseIssue(level="info", code="SEMANTIC_MAPPING", message=note, sheet="CSV", row=1))
    if not mapping and transform_mode == "standard":
        issues.append(ParseIssue(level="warning", code="FIELD_MAPPING_LOW", message="CSV表头未匹配到标准字段，请检查模板或文件命名"))
    std = normalize_dataframe(
        transformed_df, mapping, source_file, "CSV", 1,
        detected_metric=semantic.metric,
        detected_dimension=effective_dimension,
        table_type=inferred_table_type or semantic.table_type,
        semantic_confidence=semantic.confidence,
    )
    dimension_integrity = build_dimension_integrity_profile(df, std, mapping)
    validation = _template_validation(df, mapping, field_confidences, semantic, effective_dimension, dimension_integrity)
    if validation["score"] != 1.0:
        issues.append(ParseIssue(level="warning", code="TEMPLATE_VALIDATION_REVIEW", message="CSV固定模板确定性校验未全部通过；未通过项不会显示为100%。"))
    meta = [SheetMeta(
        source_file=source_file,
        sheet_name="CSV",
        sheet_type=inferred_table_type or infer_sheet_type(list(df.columns), rows, semantic.table_type),
        detected_metric=semantic.metric,
        detected_dimension=effective_dimension,
        semantic_confidence=semantic.confidence,
        dimension_confidence=effective_dimension_conf,
        field_mapping_confidence=mapping_conf,
        field_confidences=field_confidences,
        metric_validation_score=validation["metric_score"],
        dimension_validation_score=validation["dimension_score"],
        period_validation_score=validation["period_score"],
        value_validation_score=validation["value_score"],
        template_validation_score=validation["score"],
        template_validation_status=validation["status"],
        validation_details=validation,
        unmapped_named_columns=validation["field"].get("unmapped_named_columns", []),
        unlabeled_data_columns=validation["field"].get("unlabeled_data_columns", []),
        dimension_integrity=dimension_integrity,
        semantic_evidence=semantic.evidence,
        value_column=value_column,
        transform_mode=transform_mode,
        header_row=1,
        original_headers=list(df.columns),
        field_mapping=mapping,
        rows_detected=len(std),
        physical_rows=len(df) + 1,
        physical_columns=len(df.columns),
        source_data_rows=len(df),
        period_column_count=period_profile["period_column_count"],
        period_start=period_profile["period_start"],
        period_end=period_profile["period_end"],
        period_cells_total=period_profile["period_cells_total"],
        period_valid_cells=period_profile["period_valid_cells"],
        period_placeholder_cells=period_profile["period_placeholder_cells"],
        period_placeholder_breakdown=period_profile["period_placeholder_breakdown"],
    )]
    issues.append(ParseIssue(level="info", code="CSV_DIALECT", message=f"已识别CSV编码/分隔符：分隔符={repr(sep)}"))
    return std, meta, issues


def parse_file(path: Path, source_name: str | None = None) -> tuple[pd.DataFrame, list[SheetMeta], list[ParseIssue]]:
    suffix = path.suffix.lower()
    if suffix in [".xlsx", ".xlsm"]:
        return parse_excel(path, source_name=source_name)
    if suffix == ".csv":
        return parse_csv(path, source_name=source_name)
    raise ValueError("仅支持 .xlsx/.xlsm/.csv 文件")

from __future__ import annotations

from collections import Counter
import re
from typing import Any

import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models.schemas import ParseIssue, SheetMeta
from app.services.template_registry import clean_token, exact_field_for_header
from app.services.value_parser import normalize_text, parse_numeric
from .template_classifier import TemplateDecision, MONTH_RE, YEAR_RE

SPECIAL_HEADER_ALIASES = {
    clean_token("企业简称"): "oem",
    clean_token("企业"): "oem",
    clean_token("中国整车集团"): "oem",
    clean_token("整车厂"): "manufacturer",
    clean_token("国产进口"): "origin_type",
    clean_token("车辆类型"): "vehicle_type",
    clean_token("HV/PHV/EV"): "power_type",
}
IGNORE_HEADERS = {clean_token(x) for x in ["调整值", "年度", "年合计", "年度合计", "备注", "说明", "序号"]}


def _norm(v: Any) -> str:
    return normalize_text(v).strip()


def _is_month(v: Any) -> int | None:
    text = _norm(v).replace(" ", "")
    if not MONTH_RE.fullmatch(text):
        return None
    return int(text[:-1])


def _year(v: Any) -> int | None:
    text = _norm(v).replace(" ", "")
    m = YEAR_RE.fullmatch(text)
    return int(m.group(1)) if m else None


def _alias(v: Any, decision: TemplateDecision) -> str | None:
    text = _norm(v)
    token = clean_token(text)
    if not token:
        return None
    if token in IGNORE_HEADERS or _is_month(text) or _year(text):
        return "__ignore__"
    if token in SPECIAL_HEADER_ALIASES:
        return SPECIAL_HEADER_ALIASES[token]
    # Export annual template uses a column literally named EV whose cell values are EV/PHV.
    if decision.metric == "export" and token == clean_token("EV"):
        return "power_type"
    return exact_field_for_header(text)


def _is_summary(values: list[Any], dim_indices: list[int]) -> bool:
    texts = [_norm(values[i]) for i in dim_indices if i < len(values) and _norm(values[i])]
    for text in texts:
        compact = re.sub(r"\s+", "", text).lower()
        if compact in {"total", "grandtotal", "总计", "合计", "小计", "全部"}:
            return True
        if compact.endswith("total") or compact.endswith("合计") or compact.endswith("总计"):
            return True
    return False


def _unique(series: pd.Series) -> set[str]:
    out: set[str] = set()
    for x in series.tolist():
        if x is None or (isinstance(x, float) and pd.isna(x)):
            continue
        t = _norm(x)
        if t:
            out.add(t)
    return out


def parse_annual_month_blocks(
    ws: Worksheet,
    source_file: str,
    decision: TemplateDecision,
    merged_ranges: list[str] | None = None,
) -> tuple[pd.DataFrame, SheetMeta, list[ParseIssue]]:
    issues: list[ParseIssue] = []
    header_rows = sorted(decision.header_rows)
    all_records: list[dict[str, Any]] = []
    field_mapping: dict[str, str] = {}
    field_confidences: dict[str, float] = {}
    unknown_headers: set[str] = set()
    original_headers: list[str] = []
    placeholder_counter: Counter[str] = Counter()
    valid_cells = 0
    detail_row_count = 0
    periods: list[str] = []
    source_dimension_values: dict[str, set[str]] = {}

    for block_idx, header_row in enumerate(header_rows):
        next_header = header_rows[block_idx + 1] if block_idx + 1 < len(header_rows) else ws.max_row + 1
        headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
        if not original_headers:
            original_headers = [_norm(x) or f"列{i+1}" for i, x in enumerate(headers)]
        year = next((_year(x) for x in headers if _year(x) is not None), None)
        month_cols: dict[int, int] = {}
        dim_cols: dict[int, str] = {}
        for idx, h in enumerate(headers):
            month = _is_month(h)
            if month:
                month_cols[idx] = month
                continue
            a = _alias(h, decision)
            if a == "__ignore__" or not _norm(h):
                continue
            if a:
                dim_cols[idx] = a
                field_mapping[_norm(h)] = a
                field_confidences[_norm(h)] = 1.0
            else:
                # Named columns outside months/year/known ignores must be explicitly reviewed.
                unknown_headers.add(_norm(h))
        if year is None or len(month_cols) < 6:
            issues.append(ParseIssue(level="error", code="ANNUAL_BLOCK_PERIOD_ERROR", sheet=ws.title, row=header_row, message="年度分块表未找到明确年份或足够月份列"))
            continue

        carry: dict[str, Any] = {}
        fill_dims = {"oem", "manufacturer", "brand", "origin_type", "vehicle_type", "region", "size_class"}
        dim_indices = list(dim_cols.keys())
        for r in range(header_row + 1, next_header):
            values = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            # Ignore completely empty rows.
            if not any(_norm(x) for x in values):
                continue
            if _is_summary(values, dim_indices):
                continue
            dims: dict[str, Any] = {}
            for idx, field in dim_cols.items():
                value = values[idx] if idx < len(values) else None
                if _norm(value):
                    carry[field] = value
                    dims[field] = value
                elif field in fill_dims and field in carry:
                    dims[field] = carry[field]
                else:
                    dims[field] = None
            numeric_months = 0
            for idx, month in month_cols.items():
                raw = values[idx] if idx < len(values) else None
                num = parse_numeric(raw)
                if num is None:
                    text = _norm(raw)
                    placeholder_counter[text if text else "(空白)"] += 1
                    continue
                numeric_months += 1
                valid_cells += 1
                period = f"{year:04d}-{month:02d}"
                periods.append(period)
                rec = dict(dims)
                rec.update({
                    "time_period": period,
                    decision.metric: float(num),
                    "source_file": source_file,
                    "source_sheet": ws.title,
                    "source_sheet_name": ws.title,
                    "source_row": r,
                    "source_column": _norm(headers[idx]) or period,
                    "source_column_index": idx + 1,
                    "source_cell": f"{get_column_letter(idx + 1)}{r}",
                    "source_metric_type": decision.metric,
                    "source_metric_scope": decision.metric_scope,
                    "source_metric_label": decision.metric_label,
                    "source_dimension": "multi_dimension" if len({x for x in dim_cols.values() if x}) >= 2 else (next(iter(dim_cols.values())) if dim_cols else None),
                    "source_table_type": "multi_dimension_detail",
                    "source_parser_id": decision.parser_id,
                    "source_semantic_confidence": 1.0,
                    "source_row_kind": "detail",
                })
                all_records.append(rec)
            if numeric_months:
                detail_row_count += 1
                for field, val in dims.items():
                    if _norm(val):
                        source_dimension_values.setdefault(field, set()).add(_norm(val))

    std = pd.DataFrame(all_records)
    mapped_dims = [x for x in field_mapping.values() if x not in {"__ignore__"}]
    mapped_dim_set = set(mapped_dims)
    dimension_integrity: dict[str, Any] = {}
    dimension_ok = True
    for field, source_vals in source_dimension_values.items():
        if field not in std.columns:
            missing = sorted(source_vals)
        else:
            missing = sorted(source_vals - _unique(std[field]))
        dimension_integrity[field] = {
            "source_unique_count": len(source_vals),
            "standardized_unique_count": len(_unique(std[field])) if field in std.columns else 0,
            "missing_after_standardization": missing[:30],
            "status": "完整" if not missing else "有缺失",
        }
        if missing:
            dimension_ok = False

    metric_score = 1.0 if decision.metric else 0.0
    field_score = 1.0 if field_mapping and not unknown_headers else 0.0
    period_score = 1.0 if header_rows and periods and all(re.fullmatch(r"20\d{2}-\d{2}", p) for p in periods) else 0.0
    value_score = 1.0 if valid_cells > 0 and not std.empty else 0.0
    dimension_score = 1.0 if mapped_dim_set and dimension_ok else 0.0
    template_score = 1.0 if all(x == 1.0 for x in [metric_score, field_score, period_score, value_score, dimension_score]) else 0.0
    status = "PASS" if template_score == 1.0 else "REVIEW"

    if template_score == 1.0:
        issues.append(ParseIssue(level="info", code="ANNUAL_BLOCK_TEMPLATE_PASS", sheet=ws.title, message=f"年度分块模板确定性校验100%，已转换{len(std)}条月度记录"))
    else:
        issues.append(ParseIssue(level="warning", code="ANNUAL_BLOCK_TEMPLATE_REVIEW", sheet=ws.title, message=f"年度分块模板未全部通过：未知字段={sorted(unknown_headers)}"))

    periods = sorted(set(periods))
    meta = SheetMeta(
        sheet_name=ws.title,
        sheet_type="multi_dimension_detail",
        source_file=source_file,
        detected_metric=decision.metric,
        detected_dimension="multi_dimension" if len(mapped_dim_set) >= 2 else (next(iter(mapped_dim_set)) if mapped_dim_set else None),
        semantic_confidence=metric_score,
        dimension_confidence=dimension_score,
        field_mapping_confidence=field_score,
        field_confidences=field_confidences,
        metric_validation_score=metric_score,
        dimension_validation_score=dimension_score,
        period_validation_score=period_score,
        value_validation_score=value_score,
        template_validation_score=template_score,
        template_validation_status=status,
        validation_details={
            "parser": decision.parser_id,
            "structure": decision.structure,
            "metric": decision.metric,
            "scope": decision.metric_scope,
            "block_header_rows": header_rows,
            "unknown_named_columns": sorted(unknown_headers),
            "field_check": "100%" if field_score else "REVIEW",
            "period_check": "100%" if period_score else "REVIEW",
            "value_check": "100%" if value_score else "REVIEW",
            "dimension_check": "100%" if dimension_score else "REVIEW",
        },
        unmapped_named_columns=sorted(unknown_headers),
        unlabeled_data_columns=[],
        dimension_integrity=dimension_integrity,
        semantic_evidence=decision.evidence,
        value_column=None,
        transform_mode="annual_blocks_to_long",
        header_row=header_rows[0] if header_rows else None,
        original_headers=original_headers,
        field_mapping=field_mapping,
        rows_detected=len(std),
        physical_rows=ws.max_row,
        physical_columns=ws.max_column,
        source_data_rows=detail_row_count,
        period_column_count=len(periods),
        period_start=periods[0] if periods else None,
        period_end=periods[-1] if periods else None,
        period_cells_total=valid_cells + sum(placeholder_counter.values()),
        period_valid_cells=valid_cells,
        period_placeholder_cells=sum(placeholder_counter.values()),
        period_placeholder_breakdown=dict(placeholder_counter.most_common(10)),
        merged_ranges=merged_ranges or [],
        parser_id=decision.parser_id,
        metric_scope=decision.metric_scope,
        metric_label=decision.metric_label,
        structure_type=decision.structure,
        block_count=len(header_rows),
    )
    return std, meta, issues

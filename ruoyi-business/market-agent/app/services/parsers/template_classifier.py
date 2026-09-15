from __future__ import annotations

from dataclasses import dataclass, field
import re
from typing import Any

from openpyxl.worksheet.worksheet import Worksheet

from app.services.template_registry import clean_token, exact_field_for_header

YYYYMM_RE = re.compile(r"^20\d{2}(?:0[1-9]|1[0-2])(?:\.0)?$")
MONTH_RE = re.compile(r"^(?:0?[1-9]|1[0-2])月$")
YEAR_RE = re.compile(r"^(20\d{2})年?$")


@dataclass
class TemplateDecision:
    parser_id: str
    metric: str | None
    metric_scope: str | None = None
    metric_label: str | None = None
    status: str = "REVIEW"
    confidence: float = 0.0
    structure: str = "unknown"
    evidence: list[str] = field(default_factory=list)
    header_rows: list[int] = field(default_factory=list)


def _text(v: Any) -> str:
    return "" if v is None else str(v).strip()


def _row_values(ws: Worksheet, row: int) -> list[Any]:
    return [ws.cell(row, c).value for c in range(1, ws.max_column + 1)]


def _metric_from_text(text: str, source_label: str) -> tuple[str | None, str | None, str | None, list[str]]:
    token = clean_token(text)
    evidence: list[str] = []
    if not token:
        return None, None, None, evidence
    if ("出口" in token or "export" in token or "overseas" in token) and ("批发" in token or "wholesale" in token):
        return "export", "wholesale", "出口批发销量", [f"{source_label}明确包含出口+批发"]
    if ("出口" in token or "export" in token) and ("销量" in token or "sales" in token):
        return "export", "export", "出口销量", [f"{source_label}明确包含出口销量"]
    if any(x in token for x in ["国内批发", "中国市场批发", "domesticwholesale"]):
        return "domestic_wholesale", "domestic", "国内批发销量", [f"{source_label}明确包含国内/中国市场批发"]
    if "零售" in token or "retail" in token:
        if any(x in token for x in ["中国市场", "国内", "china"]):
            return "retail_sales", "domestic", "中国市场零售销量", [f"{source_label}明确包含中国/国内+零售"]
        if any(x in token for x in ["其他国家", "海外", "国外", "foreign", "overseas"]):
            return "retail_sales", "foreign", "其他国家零售销量", [f"{source_label}明确包含其他国家/海外+零售"]
        return "retail_sales", "retail", "零售销量", [f"{source_label}明确包含零售销量"]
    if "批发" in token or "wholesale" in token:
        return "wholesale", "total" if "总" in token else "wholesale", "批发销量", [f"{source_label}明确包含批发销量"]
    if any(x in token for x in ["产量", "生产量", "production", "output"]):
        return "production", None, "产量", [f"{source_label}明确包含产量"]
    if any(x in token for x in ["库存", "inventory", "stock"]):
        return "inventory", None, "库存", [f"{source_label}明确包含库存"]
    if any(x in token for x in ["国内销量", "内销量", "domesticsales"]):
        return "domestic_sales", "domestic", "国内销量", [f"{source_label}明确包含国内销量"]
    if any(x in token for x in ["销量", "销售量", "sales"]):
        return "sales", None, "销量", [f"{source_label}明确包含销量"]
    return None, None, None, evidence


def _explicit_metric(file_name: str, sheet_name: str):
    # Sheet name has higher authority than file name for a multi-Sheet workbook.
    sm, ss, sl, se = _metric_from_text(sheet_name, "Sheet名")
    if sm:
        return sm, ss, sl, se
    return _metric_from_text(file_name, "文件名")


def _header_field_count(values: list[Any]) -> int:
    return sum(1 for v in values if exact_field_for_header(v))


def _structure(ws: Worksheet) -> tuple[str, list[int], list[str]]:
    wide_candidates: list[tuple[int, int, int]] = []
    annual_headers: list[int] = []
    evidence: list[str] = []
    for r in range(1, min(ws.max_row, 120) + 1):
        vals = _row_values(ws, r)
        yyyymm = sum(1 for v in vals if YYYYMM_RE.fullmatch(_text(v).replace(" ", "")))
        months = sum(1 for v in vals if MONTH_RE.fullmatch(_text(v).replace(" ", "")))
        years = sum(1 for v in vals if YEAR_RE.fullmatch(_text(v).replace(" ", "")))
        fields = _header_field_count(vals)
        if yyyymm >= 2:
            wide_candidates.append((yyyymm, fields, r))
        if months >= 6 and years >= 1:
            annual_headers.append(r)
    if annual_headers:
        evidence.append(f"检测到{len(annual_headers)}个‘年份+1~12月’数据块表头：{annual_headers}")
        return "annual_month_blocks", annual_headers, evidence
    if wide_candidates:
        wide_candidates.sort(reverse=True)
        row = wide_candidates[0][2]
        evidence.append(f"检测到YYYYMM宽表，表头行为第{row}行")
        return "yyyymm_wide", [row], evidence
    best: tuple[int, int] | None = None
    for r in range(1, min(ws.max_row, 100) + 1):
        vals = _row_values(ws, r)
        score = _header_field_count(vals)
        if score >= 2 and (best is None or score > best[0]):
            best = (score, r)
    if best:
        evidence.append(f"检测到标准长表候选表头第{best[1]}行")
        return "long_table", [best[1]], evidence
    return "unknown", [], evidence


def _annual_structural_metric(ws: Worksheet, header_rows: list[int]):
    """Recognize registered annual-block templates from structure rather than names.

    This lets renamed files/Sheets continue to route correctly. Fingerprints are
    intentionally strict; a genuinely new/ambiguous structure remains REVIEW.
    """
    tokens: set[str] = set()
    for r in header_rows[:4]:
        for v in _row_values(ws, r):
            t = _text(v)
            if t:
                tokens.add(clean_token(t))
    domestic_signals = {clean_token("国产进口"), clean_token("车辆类型"), clean_token("HV/PHV/EV")}
    if len(tokens & domestic_signals) >= 2 and any(clean_token(x) in tokens for x in ["品牌", "车型", "企业简称"]):
        return "retail_sales", "domestic", "中国市场零售销量", ["表头结构命中已注册‘国内零售年度分块’指纹"]
    export_core = sum(1 for x in ["中国整车集团", "整车厂", "品牌", "车型"] if clean_token(x) in tokens)
    export_aux = any(clean_token(x) in tokens for x in ["调整值", "ev"])
    if export_core >= 3 and export_aux and clean_token("车辆类型") not in tokens:
        return "export", "wholesale", "出口批发销量", ["表头结构命中已注册‘出口批发年度分块’指纹"]
    return None, None, None, []


def classify_template(file_name: str, sheet_name: str, ws: Worksheet) -> TemplateDecision:
    structure, header_rows, evidence = _structure(ws)
    metric, scope, label, metric_evidence = _explicit_metric(file_name, sheet_name)
    evidence.extend(metric_evidence)

    # Structure overrides a generic “sales/销量” name when the registered template
    # contains enough business fields to distinguish domestic retail vs export.
    if structure == "annual_month_blocks":
        sm, ss, sl, se = _annual_structural_metric(ws, header_rows)
        if sm and (metric is None or metric == "sales"):
            metric, scope, label = sm, ss, sl
            evidence.extend(se)

    # Narrow registered fingerprint for the client's multi-dimensional YYYYMM
    # automotive monthly detail table. It does not apply to arbitrary numeric sheets.
    if metric is None and structure == "yyyymm_wide" and header_rows:
        vals = _row_values(ws, header_rows[0])
        mapped = {exact_field_for_header(v) for v in vals if exact_field_for_header(v)}
        automotive_dims = {"region", "oem", "brand", "vehicle_type", "model", "power_type"}
        if len(mapped & automotive_dims) >= 4:
            metric, scope, label = "sales", "market_detail", "销量"
            evidence.append("名称未写指标，但命中甲方注册‘多维汽车月度销量明细+YYYYMM’结构指纹")

    if metric is None and structure == "long_table" and header_rows:
        vals = _row_values(ws, header_rows[0])
        metric_fields = [exact_field_for_header(v) for v in vals]
        metric_fields = [x for x in metric_fields if x in {"production","sales","retail_sales","wholesale","domestic_sales","domestic_wholesale","export","inventory"}]
        if len(set(metric_fields)) == 1:
            metric = metric_fields[0]
            label = metric
            evidence.append("表头含唯一确定指标字段")

    status = "PASS" if structure != "unknown" and metric is not None else "REVIEW"
    return TemplateDecision(
        parser_id={"yyyymm_wide": "yyyymm_wide_parser", "annual_month_blocks": "annual_month_blocks_parser", "long_table": "registered_long_table_parser"}.get(structure, "unknown_parser"),
        metric=metric,
        metric_scope=scope,
        metric_label=label,
        status=status,
        confidence=1.0 if status == "PASS" else 0.0,
        structure=structure,
        evidence=evidence,
        header_rows=header_rows,
    )

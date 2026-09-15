from pathlib import Path
from openpyxl import Workbook

from app.services.parser import parse_file
from app.services.market_analysis import MarketAnalyzer


def _add_wide_sheet(wb: Workbook, title: str, dim_header: str = "名称"):
    if wb.worksheets and wb.worksheets[0].title == "Sheet" and wb.worksheets[0].max_row == 1 and wb.worksheets[0].max_column == 1:
        ws = wb.worksheets[0]
        ws.title = title
    else:
        ws = wb.create_sheet(title)
    ws.append([dim_header, 202401, 202501])
    ws.append(["对象A", 100, 120])
    ws.append(["对象B", 80, 90])
    return ws


def test_all_registered_metric_sheet_types_are_deterministic_100(tmp_path: Path):
    cases = [
        ("市场销量", "sales", "指标"),
        ("市场零售销量", "retail_sales", "指标"),
        ("市场批发销量", "wholesale", "指标"),
        ("市场产量", "production", "指标"),
        ("国内销量", "domestic_sales", "指标"),
        ("国内批发销量", "domestic_wholesale", "指标"),
        ("出口", "export", "指标"),
        ("库存", "inventory", "指标"),
        ("OEM销量", "sales", "名称"),
        ("品牌销量", "sales", "规格"),
        ("车型销量", "sales", "规格"),
        ("动力类型销量", "sales", "类型"),
        ("车种销量", "sales", "名称"),
        ("国家销量", "sales", "名称"),
        ("级别销量", "sales", "名称"),
    ]
    path = tmp_path / "固定模板全集.xlsx"
    wb = Workbook()
    wb.remove(wb.active)
    for title, _, dim_header in cases:
        ws = wb.create_sheet(title)
        ws.append([dim_header, 202401, 202501])
        ws.append(["对象A", 100, 120])
        ws.append(["对象B", 80, 90])
    wb.save(path)

    df, metas, issues = parse_file(path, source_name=path.name)
    assert len(metas) == len(cases)
    by_sheet = {m.sheet_name: m for m in metas}
    for title, metric, _ in cases:
        m = by_sheet[title]
        assert m.detected_metric == metric, (title, m.detected_metric)
        assert m.semantic_confidence == 1.0
        assert m.field_mapping_confidence == 1.0
        assert m.metric_validation_score == 1.0
        assert m.dimension_validation_score == 1.0
        assert m.period_validation_score == 1.0
        assert m.value_validation_score == 1.0
        assert m.template_validation_score == 1.0
        assert m.template_validation_status == "PASS"


def test_specific_sales_scopes_are_not_renamed(tmp_path: Path):
    path = tmp_path / "汽车销量数据.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "总零售销量"
    ws.append(["车型", 202401, 202501])
    ws.append(["车型A", 100, 120])
    ws2 = wb.create_sheet("批发销量")
    ws2.append(["车型", 202401, 202501])
    ws2.append(["车型A", 130, 150])
    wb.save(path)
    df, metas, _ = parse_file(path, source_name=path.name)
    assert metas[0].detected_metric == "retail_sales"
    assert metas[0].template_validation_score == 1.0
    assert metas[1].detected_metric == "wholesale"
    assert metas[1].template_validation_score == 1.0
    assert "retail_sales" in df.columns
    assert "wholesale" in df.columns


def test_unknown_or_conflicting_template_is_not_faked_to_100(tmp_path: Path):
    path = tmp_path / "未知数据.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["神秘维度", 202401, 202501])
    ws.append(["A", 1, 2])
    wb.save(path)
    _, metas, issues = parse_file(path, source_name=path.name)
    m = metas[0]
    assert m.template_validation_score == 0.0
    assert m.template_validation_status == "REVIEW"
    assert m.semantic_confidence == 0.0
    assert any(i.code == "TEMPLATE_VALIDATION_REVIEW" for i in issues)


def test_actual_toyota_retail_file_core_template_passes_100_if_available():
    path = Path("/mnt/data/海外整车厂--丰田2023、2024年24个月月度汽车销量.xlsx")
    if not path.exists():
        return
    df, metas, _ = parse_file(path, source_name=path.name)
    m = metas[0]
    assert m.detected_metric == "retail_sales"
    assert m.template_validation_score == 1.0
    assert m.template_validation_status == "CORE_PASS_AUX_REVIEW"
    assert m.unlabeled_data_columns == ["列32", "列33"]
    analyzer = MarketAnalyzer(df)
    assert analyzer.primary_sales_metric() == "retail_sales"
    result = analyzer.run_all()
    assert result.power_share
    assert abs(sum(float(x["占比"]) for x in result.power_share) - 1.0) < 1e-9

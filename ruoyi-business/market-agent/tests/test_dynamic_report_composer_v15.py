from __future__ import annotations

import json
import shutil
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from app.services.chat_engine import ChatEngine
from app.services.market_components import MarketComponentEngine, line_chart_component_id
from app.services.parser import parse_file
from app.services.report_generator import build_report_payload, export_report
from app.services.report_planner import apply_report_plan_instruction, _rule_patch
from app.services.storage import processed_base, save_dataset


ROOT = Path(__file__).resolve().parents[1]


def _timeseries_df():
    df, _, _ = parse_file(ROOT / "data" / "sample_timeseries_market.xlsx")
    return df


def test_specific_dashboard_chart_is_a_first_class_component():
    df = _timeseries_df()
    capabilities = MarketComponentEngine(df).capabilities()
    patch = _rule_patch("把动力类型月度结构占比图加入周报", capabilities)
    assert line_chart_component_id("动力类型月度结构占比") in patch["components"]
    assert "power_structure" not in patch["components"]


def test_current_selected_chart_can_be_added_by_reference():
    df = _timeseries_df()
    capabilities = MarketComponentEngine(df).capabilities()
    patch = _rule_patch("把这张图加入周报", capabilities, selected_chart_title="动力类型月度结构占比")
    assert patch["components"] == [line_chart_component_id("动力类型月度结构占比")]


def test_explicit_title_wins_and_is_not_treated_as_a_filter():
    df = _timeseries_df()
    capabilities = MarketComponentEngine(df).capabilities()
    patch = _rule_patch(
        "再新增一个2024年市场观察，标题为新能源结构专题，并加入动力类型月度销量",
        capabilities,
    )
    assert patch["custom_title"] == "新能源结构专题"
    assert patch["filters"] == {}


def test_title_month_updates_observation_period_and_payload_is_json_safe():
    df = _timeseries_df()
    dataset_id = "pytest_dynamic_composer_period"
    try:
        save_dataset(dataset_id, df, {"file_name": "sample_timeseries_market.xlsx"})
        plan, _ = apply_report_plan_instruction(
            dataset_id,
            "生成美国市场11月观察，加入动力类型月度结构占比图",
            df,
            use_llm=False,
        )
        observation = plan["market_observations"][0]
        assert observation["custom_title"] == "美国市场11月观察"
        assert observation["analysis_period"]["start_period"] == "2025-11"
        assert line_chart_component_id("动力类型月度结构占比") in observation["components"]
        payload = build_report_payload(df, use_llm=False, report_plan=plan)
        assert payload["market_observations"][0]["period_label"] == "2025年11月"
        json.dumps(payload, ensure_ascii=False, allow_nan=False)
    finally:
        shutil.rmtree(processed_base(dataset_id), ignore_errors=True)


def test_multiple_observations_and_no_observation_mode():
    df = _timeseries_df()
    dataset_id = "pytest_dynamic_composer_multiple"
    try:
        save_dataset(dataset_id, df, {"file_name": "sample_timeseries_market.xlsx"})
        first, _ = apply_report_plan_instruction(dataset_id, "生成美国市场11月观察，放总体市场", df, use_llm=False)
        assert len(first["market_observations"]) == 1
        second, _ = apply_report_plan_instruction(dataset_id, "再新增一个新能源汽车市场观察，加入动力类型月度结构占比图", df, use_llm=False)
        assert len(second["market_observations"]) == 2
        assert second["market_observations"][1]["custom_title"] == "新能源汽车市场观察"
        removed, _ = apply_report_plan_instruction(dataset_id, "本期不要市场观察", df, use_llm=False)
        assert removed["market_observations"] == []
        payload = build_report_payload(df, use_llm=False, report_plan=removed)
        assert payload["dynamic_report_plan"] is True
        assert payload["market_observations"] == []
    finally:
        shutil.rmtree(processed_base(dataset_id), ignore_errors=True)


def test_chat_uses_selected_dashboard_chart_context():
    df = _timeseries_df()
    dataset_id = "pytest_dynamic_composer_chat"
    try:
        save_dataset(dataset_id, df, {"file_name": "sample_timeseries_market.xlsx"})
        response = ChatEngine(df, dataset_id=dataset_id).answer(
            "帮我把这张图加入到周报中",
            use_llm=False,
            selected_chart_title="动力类型月度结构占比",
        )
        assert response.report_updated is True
        components = response.report_plan["market_observation"]["components"]
        assert line_chart_component_id("动力类型月度结构占比") in components
    finally:
        shutil.rmtree(processed_base(dataset_id), ignore_errors=True)


def test_multiple_observations_export_to_all_formats(tmp_path):
    df = _timeseries_df()
    plan = {
        "mode": "custom",
        "market_observations": [
            {
                "id": "market_observation_1",
                "custom_title": "美国11月市场观察",
                "components": ["market_summary", line_chart_component_id("动力类型月度结构占比")],
                "summary_segments": ["total_market", "new_energy_vehicle"],
                "analysis_period": {"period_mode": "single", "start_period": "2025-11", "end_period": "2025-11"},
            },
            {
                "id": "market_observation_2",
                "custom_title": "新能源汽车趋势观察",
                "components": [line_chart_component_id("新能源汽车市场占有率")],
            },
        ],
        "active_observation_id": "market_observation_2",
    }
    paths = {fmt: export_report(df, fmt, tmp_path, "v15", report_plan=plan) for fmt in ["xlsx", "docx", "pptx"]}
    assert len(load_workbook(paths["xlsx"], read_only=True).sheetnames) >= 5
    assert any("美国11月市场观察" in paragraph.text for paragraph in Document(paths["docx"]).paragraphs)
    presentation = Presentation(paths["pptx"])
    slide_text = "\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text"))
    assert "美国11月市场观察" in slide_text
    assert "新能源汽车趋势观察" in slide_text

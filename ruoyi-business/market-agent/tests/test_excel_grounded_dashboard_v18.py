from __future__ import annotations

import math

import pandas as pd
from openpyxl import Workbook

from app.services.market_analysis import MarketAnalyzer
from app.services.report_generator import write_table


def _source_df() -> pd.DataFrame:
    rows = []
    source_row = 2
    for year in [2024, 2025]:
        for month in range(1, 13):
            period = f"{year}-{month:02d}"
            values = [
                ("通用集团", "ICE", 200_000 + month * 1_000),
                ("丰田集团", "HV", 150_000 + month * 800),
                ("福特集团", "EV", 100_000 + month * 600),
            ]
            for oem, power, sales in values:
                rows.append({
                    "time_period": period, "region": "美国", "oem": oem,
                    "brand": oem.replace("集团", ""), "model": oem + "车型", "power_type": power,
                    "sales": float(sales), "source_file": "美国市场月度销量.xlsx",
                    "source_sheet": "OEM销量明细", "source_row": source_row,
                    "source_metric_type": "sales", "source_metric_label": "销量",
                    "source_dimension": "oem/power_type",
                })
                source_row += 1
    return pd.DataFrame(rows)


def _chart(result, title: str) -> dict:
    return next(x for x in result.line_charts if x["title"] == title)


def test_latest_and_same_single_month_have_identical_trends():
    df = _source_df()
    latest = MarketAnalyzer(df, period_mode="latest").run_all()
    single = MarketAnalyzer(df, period_mode="single", start_period="2025-12").run_all()
    assert latest.summary == single.summary
    assert latest.monthly_trend == single.monthly_trend
    assert latest.line_charts == single.line_charts
    assert latest.analysis_period["trend_available_periods"] == single.analysis_period["trend_available_periods"]
    assert _chart(single, "动力类型月度结构占比")["series"] == _chart(latest, "动力类型月度结构占比")["series"]


def test_single_earlier_month_trend_is_cut_off_without_future_data():
    result = MarketAnalyzer(_source_df(), period_mode="single", start_period="2025-09").run_all()
    assert result.monthly_trend[-1]["时间"] == "2025-09"
    for chart in result.line_charts:
        assert all(point["时间"] <= "2025-09" for series in chart["series"] for point in series["points"])
    assert result.analysis_period["trend_scope_label"] == "截至2025年9月的全部可用月份"


def test_ranking_power_and_share_reconcile_to_uploaded_rows():
    df = _source_df()
    target = df[df["time_period"] == "2025-12"].copy()
    result = MarketAnalyzer(df, period_mode="single", start_period="2025-12").run_all()
    manual_oem = target.groupby("oem", as_index=False)["sales"].sum().sort_values("sales", ascending=False)
    actual = result.rankings["oem"]
    denominator = float(target["sales"].sum())
    assert [x["对象"] for x in actual] == manual_oem["oem"].tolist()
    assert [x["销量/数值"] for x in actual] == manual_oem["sales"].tolist()
    for row in actual:
        assert math.isclose(row["占比"], row["销量/数值"] / denominator, rel_tol=0, abs_tol=1e-12)
    manual_power = target.groupby("power_type", as_index=False)["sales"].sum().sort_values("sales", ascending=False)
    power = result.power_share
    assert [x["动力类型"] for x in power] == manual_power["power_type"].tolist()
    assert [x["销量/数值"] for x in power] == manual_power["sales"].tolist()
    assert math.isclose(sum(x["占比"] for x in power), 1.0, rel_tol=0, abs_tol=1e-12)


def test_cached_rankings_and_charts_are_equivalent_and_return_safe_copies():
    analyzer = MarketAnalyzer(_source_df(), period_mode="single", start_period="2025-12")
    full_ranking = analyzer.ranking("oem", top_n=None)
    top_two = analyzer.ranking("oem", top_n=2)
    assert top_two == full_ranking[:2]
    top_two[0]["对象"] = "被调用方修改"
    assert analyzer.ranking("oem", top_n=1)[0]["对象"] == full_ranking[0]["对象"]

    first_charts = analyzer.line_charts()
    assert first_charts
    expected_title = first_charts[0]["title"]
    first_charts[0]["title"] = "被调用方修改"
    assert analyzer.line_charts()[0]["title"] == expected_title


def test_calculation_audit_exposes_excel_lineage_and_disables_llm_math():
    result = MarketAnalyzer(_source_df(), period_mode="single", start_period="2025-12").run_all()
    audit = result.calculation_audit
    assert audit["计算引擎"] == "Python/Pandas确定性计算"
    assert audit["LLM参与数值计算"] is False
    assert audit["趋势周期"] == "截至2025年12月的全部可用月份"
    oem = next(x for x in audit["模块"] if x["模块ID"] == "ranking:oem")
    assert oem["源文件"] == ["美国市场月度销量.xlsx"]
    assert oem["源Sheet"] == ["OEM销量明细"]
    assert oem["标准数值字段"] == "sales"
    assert oem["参与标准化记录数"] == 3
    assert oem["占比恒等式通过"] is True


def test_excel_export_keeps_numeric_share_and_percentage_format():
    wb = Workbook()
    ws = wb.active
    write_table(ws, 1, "OEM排名", [{"排名": 1, "对象": "通用集团", "销量/数值": 263459, "占比": 0.1732}])
    assert ws.cell(3, 4).value == 0.1732
    assert ws.cell(3, 4).number_format == "0.00%"

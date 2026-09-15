from pathlib import Path
import shutil
from openpyxl import Workbook
from fastapi.testclient import TestClient

from app.main import app
from app.services.parser import parse_file
from app.services.market_analysis import MarketAnalyzer
from app.services.storage import (
    processed_base,
    save_dataset,
    save_sheet_datasets,
    load_meta,
    load_sheet_dataset,
    available_sheets,
    build_comprehensive_dataset,
)


def _add_wide(ws, title, metric_rows):
    ws.title = title
    ws.append(["集团", "整车厂/品牌", "国家/地区", "车种", "级别", "车型", "动力总成", 202301, 202302, 202401, 202402])
    for row in metric_rows:
        ws.append(row)


def _add_annual_domestic(ws, title):
    ws.title = title
    headers = ["企业简称", "品牌", "国产进口", "车辆类型", "车型", "HV/PHV/EV"] + [f"{m}月" for m in range(1,13)] + ["2023年"]
    ws.append(headers)
    ws.append(["A企","A牌","国产","Cars","车型1","EV"] + [10]*12 + [120])
    ws.append([])
    headers2 = headers[:-1] + ["2024年"]
    ws.append(headers2)
    ws.append(["A企","A牌","国产","Cars","车型1","EV"] + [20]*12 + [240])


def _add_annual_export(ws, title):
    ws.title = title
    headers = ["中国整车集团", "整车厂", "品牌", "车型", "EV"] + [f"{m}月" for m in range(1,13)] + ["调整值", "2023年"]
    ws.append(headers)
    ws.append(["A集团","A厂","A牌","车型1","EV"] + [1]*12 + [0,12])
    ws.append([])
    headers2 = headers[:-1] + ["2024年"]
    ws.append(headers2)
    ws.append(["A集团","A厂","A牌","车型1","EV"] + [2]*12 + [0,24])


def test_multisheet_independent_and_comprehensive(tmp_path: Path):
    p = tmp_path / "客户汽车月报.xlsx"
    wb = Workbook()
    _add_wide(wb.active, "总批发销量", [["A集团","A牌","中国","Cars","C","车型1","EV",100,110,200,220]])
    _add_annual_domestic(wb.create_sheet(), "中国市场零售销量")
    _add_annual_export(wb.create_sheet(), "出口批发销量")
    _add_wide(wb.create_sheet(), "其他国家零售销量", [["A集团","A牌","德国","Cars","C","车型1","EV",5,6,10,12]])
    wb.save(p)

    df, metas, issues = parse_file(p, source_name=p.name)
    assert len(metas) == 4
    assert all(m.template_validation_score == 1.0 for m in metas)
    assert {m.detected_metric for m in metas} == {"wholesale", "retail_sales", "export"}

    # Annual block sheets must really produce 24 monthly periods and be analyzable.
    for sheet in ["中国市场零售销量", "出口批发销量"]:
        part = df[df["source_sheet_name"] == sheet]
        analyzer = MarketAnalyzer(part)
        assert analyzer.latest_period() == "2024-12"
        assert len(analyzer.trend()) == 24
        assert analyzer.ranking("model")
        assert analyzer.power_structure()

    dataset_id = "v12_test_multisheet"
    save_dataset(dataset_id, df, {"file_name": p.name, "sheet_meta": [m.model_dump() for m in metas], "issues": [i.model_dump() for i in issues]})
    sheet_files = save_sheet_datasets(dataset_id, df)
    meta = load_meta(dataset_id); meta["sheet_files"] = sheet_files; save_dataset(dataset_id, df, meta)
    assert available_sheets(dataset_id)[-1] == "综合视图"
    cdf = build_comprehensive_dataset(dataset_id)
    latest = cdf[cdf["time_period"] == "2024-02"].iloc[0]
    assert latest["wholesale"] == 220
    assert latest["export"] == 2
    assert latest["retail_sales"] == 20
    assert latest["domestic_wholesale"] == 218
    assert latest["inventory"] == 198


def test_analysis_with_selected_sheet_never_uses_other_sheet_metrics(tmp_path: Path):
    """A concrete Sheet is a strict boundary for every dashboard result."""
    p = tmp_path / "分表边界.xlsx"
    wb = Workbook()
    _add_wide(wb.active, "总批发销量", [["A集团", "A牌", "中国", "Cars", "C", "车型1", "EV", 100, 110, 200, 220]])
    _add_annual_domestic(wb.create_sheet(), "中国市场零售销量")
    wb.save(p)

    df, metas, issues = parse_file(p, source_name=p.name)
    dataset_id = "pytest_selected_sheet_boundary"
    try:
        save_dataset(dataset_id, df, {
            "file_name": p.name,
            "sheet_meta": [m.model_dump() for m in metas],
            "issues": [i.model_dump() for i in issues],
        })
        sheet_files = save_sheet_datasets(dataset_id, df)
        meta = load_meta(dataset_id)
        meta["sheet_files"] = sheet_files
        save_dataset(dataset_id, df, meta)

        response = TestClient(app).get(
            f"/api/analysis/{dataset_id}",
            params={"sheet_name": "中国市场零售销量", "period_mode": "year", "year": 2024},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["analysis_scope"] == {
            "mode": "sheet",
            "sheet_name": "中国市场零售销量",
            "strict_sheet": True,
        }
        # The domestic retail sheet contains only retail_sales.  Wholesale
        # values from the first worksheet would prove an invalid fallback.
        row = body["overview_table"][0]
        assert row["指标"] == "Cars"
        assert "零售销量" in row
        assert "批发销量" not in row
        assert str(row["零售销量"]).startswith("240.0")
    finally:
        shutil.rmtree(processed_base(dataset_id), ignore_errors=True)


def test_renamed_annual_sheets_route_by_structure(tmp_path: Path):
    p = tmp_path / "任意名称.xlsx"
    wb = Workbook()
    _add_annual_domestic(wb.active, "Data_A")
    _add_annual_export(wb.create_sheet(), "Data_B")
    wb.save(p)
    _, metas, _ = parse_file(p, source_name=p.name)
    by_sheet = {m.sheet_name: m for m in metas}
    assert by_sheet["Data_A"].detected_metric == "retail_sales"
    assert by_sheet["Data_A"].metric_scope == "domestic"
    assert by_sheet["Data_A"].template_validation_score == 1.0
    assert by_sheet["Data_B"].detected_metric == "export"
    assert by_sheet["Data_B"].metric_scope == "wholesale"
    assert by_sheet["Data_B"].template_validation_score == 1.0


def test_missing_dimensions_do_not_create_fake_rankings(tmp_path: Path):
    """缺失维度不能被自动制造成“未分类(空白)”排名；已有车种应可回退到市场页签。"""
    p = tmp_path / "结构缺失维度.xlsx"
    wb = Workbook()
    _add_annual_domestic(wb.active, "Data_Domestic")
    _add_annual_export(wb.create_sheet(), "Data_Export")
    wb.save(p)
    df, metas, _ = parse_file(p, source_name=p.name)

    domestic = df[df["source_sheet_name"] == "Data_Domestic"]
    dres = MarketAnalyzer(domestic).run_all()
    # 国内零售结构没有market，但有vehicle_type=Cars，市场/车种页签应展示真实车种，而非空白伪分类。
    assert len(dres.rankings["market"]) == 1
    assert dres.rankings["market"][0]["对象"] == "Cars"
    assert "market" not in (dres.dimension_integrity or {})
    assert "vehicle_type" in (dres.dimension_integrity or {})

    export = df[df["source_sheet_name"] == "Data_Export"]
    eres = MarketAnalyzer(export).run_all()
    # 出口模板没有market/vehicle_type，不允许自动生成“未分类(空白)”市场排名。
    assert eres.rankings["market"] == []
    assert "market" not in (eres.dimension_integrity or {})
    assert "vehicle_type" not in (eres.dimension_integrity or {})
    assert eres.rankings["model"]
    assert eres.power_share

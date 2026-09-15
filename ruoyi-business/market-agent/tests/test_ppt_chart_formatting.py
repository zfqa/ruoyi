from __future__ import annotations

import zipfile

from lxml import etree
from pptx import Presentation
from pptx.util import Inches

from app.services.report_generator import add_ppt_line_chart, add_ppt_table, _add_dynamic_market_observation_slides


NS = {
    "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def _chart_xml(path, number=1):
    with zipfile.ZipFile(path) as archive:
        return etree.fromstring(archive.read(f"ppt/charts/chart{number}.xml"))


def test_compact_ppt_line_chart_uses_horizontal_sampled_month_labels(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    periods = [f"2024-{month:02d}" for month in range(1, 13)] + [f"2025-{month:02d}" for month in range(1, 13)]
    chart = {
        "title": "核心指标月度趋势",
        "y_format": "number",
        "series": [{"name": "销量", "points": [{"时间": period, "数值": 1_000_000 + index * 10_000} for index, period in enumerate(periods)]}],
    }

    add_ppt_line_chart(slide, chart, Inches(0.55), Inches(4.72), Inches(6.05), Inches(2.22))
    output = tmp_path / "compact-chart.pptx"
    presentation.save(output)

    xml = _chart_xml(output)
    assert xml.xpath("string(.//c:dateAx/c:majorUnit/@val)", namespaces=NS) == "4"
    assert xml.xpath("string(.//c:dateAx/c:majorTimeUnit/@val)", namespaces=NS) == "months"
    assert xml.xpath("string(.//c:dateAx/c:numFmt/@formatCode)", namespaces=NS) == "yy-MM"
    assert xml.xpath("string(.//c:dateAx/c:txPr/a:bodyPr/@rot)", namespaces=NS) == "0"
    assert xml.xpath("string(.//c:dateAx/c:txPr/a:bodyPr/@vert)", namespaces=NS) == "horz"
    assert xml.xpath("string(.//c:valAx/c:numFmt/@formatCode)", namespaces=NS) == '0.0,,"M"'


def test_percent_ppt_line_chart_uses_consistent_scale(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    chart = {
        "title": "动力类型月度结构占比",
        "y_format": "percent",
        "series": [{"name": "新能源汽车", "points": [{"时间": f"2025-{month:02d}", "数值": 0.1 + month / 100} for month in range(1, 13)]}],
    }

    add_ppt_line_chart(slide, chart, Inches(0.55), Inches(4.72), Inches(6.05), Inches(2.22))
    output = tmp_path / "percent-chart.pptx"
    presentation.save(output)

    xml = _chart_xml(output)
    assert xml.xpath("string(.//c:valAx/c:numFmt/@formatCode)", namespaces=NS) == "0%"
    assert xml.xpath("string(.//c:valAx/c:scaling/c:min/@val)", namespaces=NS) == "0.0"
    assert xml.xpath("string(.//c:valAx/c:scaling/c:max/@val)", namespaces=NS) == "1.0"
    assert xml.xpath("string(.//c:valAx/c:majorUnit/@val)", namespaces=NS) == "0.2"


def test_ppt_table_formats_numeric_values_for_readability():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_ppt_table(
        slide,
        [{"动力类型": "ICE", "销量/数值": 1_147_698.0, "占比": 0.7546, "销量口径": "销量"}],
        Inches(0.5), Inches(1), Inches(12), Inches(2), max_rows=None,
    )

    table = next(shape.table for shape in slide.shapes if shape.has_table)
    assert table.cell(1, 1).text == "1,147,698"
    assert table.cell(1, 2).text == "75.46%"
    assert table.rows[0].height == Inches(0.44)
    assert table.rows[1].height <= Inches(0.58)
    assert table.cell(1, 0).vertical_anchor is not None


def test_dynamic_ppt_keeps_summary_table_and_trend_charts_on_separate_slides():
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    periods = [f"2025-{month:02d}" for month in range(1, 13)]
    line_chart = {
        "title": "月度趋势", "y_format": "number",
        "series": [{"name": "销量", "points": [{"时间": period, "数值": month * 100} for month, period in enumerate(periods, 1)]}],
    }
    payload = {
        "market_observations": [{
            "number": "1.1", "title": "市场观察", "period_label": "2025年12月",
            "components": [
                {"id": "market_summary", "matrix": True, "rows": [{"指标": "总体市场", "销量": "1,000", "具体分析": "销量保持稳定。"}]},
                {"id": "power_structure", "type": "table", "title": "动力类型结构占比", "rows": [
                    {"动力类型": "ICE", "销量/数值": 800, "占比": 0.8, "销量口径": "销量"},
                    {"动力类型": "EV", "销量/数值": 200, "占比": 0.2, "销量口径": "销量"},
                ]},
                {"id": "trend_1", "type": "chart", "chart": line_chart},
                {"id": "trend_2", "type": "chart", "chart": line_chart},
            ],
        }]
    }

    end_page = _add_dynamic_market_observation_slides(presentation, blank, payload, 1)

    assert end_page == 4
    assert len(presentation.slides) == 3
    assert sum(1 for shape in presentation.slides[0].shapes if getattr(shape, "has_chart", False)) == 0
    assert sum(1 for shape in presentation.slides[1].shapes if getattr(shape, "has_chart", False)) == 1
    assert sum(1 for shape in presentation.slides[2].shapes if getattr(shape, "has_chart", False)) == 2

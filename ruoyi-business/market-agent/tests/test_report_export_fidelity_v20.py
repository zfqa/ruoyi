from __future__ import annotations

from pathlib import Path

from docx import Document
from openpyxl import load_workbook
import pandas as pd
from pptx import Presentation

from app.core.llm import LLMClient
from app.services import report_generator
from app.services.report_generator import build_report_payload, export_report
from app.services.report_planner import normalize_report_plan


def _frame():
    rows = []
    for period, factor in [("2024-11", 0.9), ("2025-10", 0.95), ("2025-11", 1.0)]:
        for index, (power, value) in enumerate([
            ("ICE", 755), ("HV", 122), ("EV", 60), ("N/A", 31),
            ("PHV", 15), ("MHV/PHV", 8), ("FCV", 2),
        ]):
            rows.append({
                "time_period": period,
                "region": "美国",
                "market": "Cars",
                "vehicle_type": "乘用车",
                "oem": f"车企{index}",
                "brand": f"品牌{index}",
                "model": f"车型{index}",
                "power_type": power,
                "sales": value * factor,
                "source_file": "sample.xlsx",
                "source_sheet": "Sheet1",
                "source_metric_type": "sales",
                "source_metric_label": "销量",
                "source_row_kind": "detail",
            })
    return pd.DataFrame(rows)


def _context_items(count: int = 9):
    items = []
    for index in range(1, count + 1):
        body = f"这是第{index}条用户资料。" + (f"完整正文{index}，" * 80) + f"尾部校验标记-{index}"
        items.append({
            "category": "personnel",
            "title": "2.1人事调整",
            "content": f"2.1人事调整\n{index}\n企业{index}\n事件标题{index}\n{body}",
            "source_name": "用户原始资料.pptx",
            "locator": f"第{index}页",
        })
    return items


def test_context_payload_keeps_every_item_and_every_character():
    payload = build_report_payload(_frame(), context_items=_context_items(), use_llm=False)

    assert len(payload["personnel"]) == 9
    assert [item["number"] for item in payload["personnel"]] == [f"2.1.{i}" for i in range(1, 10)]
    for index, item in enumerate(payload["personnel"], 1):
        assert item["title"] == f"企业{index} 事件标题{index}"
        assert f"尾部校验标记-{index}" in item["content"]
        assert item["summary"] == item["content"]
        assert "…" not in item["content"]
        assert item["analysis"] == ""


def test_legacy_filters_are_cleared_and_cannot_be_applied_invisibly():
    plan = normalize_report_plan({
        "market_observations": [{
            "id": "market_observation_1",
            "custom_title": "完整市场观察",
            "components": ["market_summary"],
            "filters": {"brand": ["旧品牌筛选"]},
        }],
    })
    assert plan["market_observations"][0]["filters"] == {}

    payload = build_report_payload(_frame(), use_llm=False, report_plan=plan)
    assert payload["market_observations"][0]["filters"] == {}
    assert payload["market_observations"][0]["components"]


def test_all_exports_preserve_context_hierarchy_and_hide_technical_tail(tmp_path, monkeypatch):
    monkeypatch.setattr(LLMClient, "enabled", property(lambda self: False))
    frame = _frame()
    context_items = _context_items()
    report_config = {
        "include_overview": False,
        "include_rankings": False,
        "include_power_structure": False,
        "include_monthly_trend": False,
        "include_line_charts": False,
        "include_weekly_content": True,
        "include_anomalies": False,
    }

    docx_path = export_report(frame, "docx", tmp_path, "fidelity", context_items=context_items, report_config=report_config)
    document = Document(docx_path)
    paragraphs = document.paragraphs
    parent_headings = [p.text for p in paragraphs if p.style.name == "Heading 2" and p.text == "2.1 人事调整"]
    child_headings = [p.text for p in paragraphs if p.style.name == "Heading 3" and p.text.startswith("2.1.")]
    full_doc_text = "\n".join(p.text for p in paragraphs)
    assert len(parent_headings) == 1
    assert len(child_headings) == 9
    for index in range(1, 10):
        assert f"尾部校验标记-{index}" in full_doc_text
    assert "提示：" not in full_doc_text
    assert "当前报告计划版本" not in full_doc_text
    assert "本报告已应用对话指令" not in full_doc_text

    xlsx_path = export_report(frame, "xlsx", tmp_path, "fidelity", context_items=context_items, report_config=report_config)
    workbook = load_workbook(xlsx_path, read_only=True)
    personnel_text = "\n".join(
        str(value)
        for row in workbook["人事调整"].iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    for index in range(1, 10):
        assert f"尾部校验标记-{index}" in personnel_text

    pptx_path = export_report(frame, "pptx", tmp_path, "fidelity", context_items=context_items, report_config=report_config)
    presentation = Presentation(pptx_path)
    slide_text = "\n".join(
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if hasattr(shape, "text")
    )
    for index in range(1, 10):
        assert f"尾部校验标记-{index}" in slide_text
        assert f"2.1.{index}" in slide_text
    assert "当前报告计划版本" not in slide_text
    assert "未发现可用于计算" not in slide_text


def test_power_chart_uses_non_overlapping_legend_layout(tmp_path):
    rows = [
        {"动力类型": "ICE", "销量/数值": 755},
        {"动力类型": "HV", "销量/数值": 122},
        {"动力类型": "EV", "销量/数值": 60},
        {"动力类型": "N/A", "销量/数值": 31},
        {"动力类型": "PHV", "销量/数值": 15},
        {"动力类型": "MHV/PHV", "销量/数值": 8},
        {"动力类型": "FCV", "销量/数值": 2},
    ]
    image = tmp_path / "power_share.png"
    assert report_generator._save_power_pie_png(rows, image)
    assert image.exists() and image.stat().st_size > 10_000

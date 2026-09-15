from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation

from app.main import app
from app.services.market_components import line_chart_component_id
from app.services.parser import parse_file
from app.services.report_config import load_report_config
from app.services.report_generator import build_report_payload, export_report
from app.services.report_planner import load_report_plan
from app.services.storage import processed_base, save_dataset


ROOT = Path(__file__).resolve().parents[1]
NEV_SHARE = line_chart_component_id("新能源汽车市场占有率")
POWER_MONTHLY = line_chart_component_id("动力类型月度结构占比")


def _frame():
    frame, _, _ = parse_file(ROOT / "data" / "sample_timeseries_market.xlsx")
    return frame


def _save(dataset_id: str):
    frame = _frame()
    save_dataset(dataset_id, frame, {"file_name": "sample_timeseries_market.xlsx"})
    return frame


def _cleanup(dataset_id: str):
    shutil.rmtree(processed_base(dataset_id), ignore_errors=True)


def test_analysis_response_embeds_visual_composer_catalog_without_second_request():
    dataset_id = "pytest_visual_catalog_in_analysis"
    try:
        _save(dataset_id)
        response = TestClient(app).get(
            f"/api/analysis/{dataset_id}",
            params={"period_mode": "single", "start_period": "2025-11"},
        )
        assert response.status_code == 200, response.text
        body = response.json()
        component_ids = {item["id"] for item in body["dashboard_components"]}
        assert "market_summary" in component_ids
        assert POWER_MONTHLY in component_ids
        assert body["dashboard_dimensions"]
    finally:
        _cleanup(dataset_id)


def test_visual_composer_persists_real_11_content_and_exports_same_order(tmp_path):
    dataset_id = "pytest_visual_report_composer"
    try:
        frame = _save(dataset_id)
        response = TestClient(app).put(
            f"/api/report-plan/{dataset_id}/visual",
            json={
                "report_title": "美国汽车市场洞察周报",
                "observation_title": "美国市场11月观察：新能源结构持续调整",
                "component_ids": ["market_summary", NEV_SHARE, POWER_MONTHLY],
                "summary_segments": [
                    "total_market", "passenger_vehicle", "commercial_vehicle", "new_energy_vehicle",
                ],
                "manual_content": "管理层补充判断：短期市场仍需关注新能源需求变化。",
                "lookback_months": 12,
                "period_mode": "single",
                "start_period": "2025-11",
                "end_period": "2025-11",
                "include_weekly_content": True,
                "include_anomalies": True,
            },
        )
        assert response.status_code == 200, response.text
        body = response.json()
        plan = body["plan"]
        observation = plan["market_observations"][0]
        assert plan["revision"] == 1
        assert plan["planner_source"] == "visual_composer"
        assert observation["custom_title"] == "美国市场11月观察：新能源结构持续调整"
        assert observation["components"] == ["market_summary", NEV_SHARE, POWER_MONTHLY]
        assert observation["analysis_period"]["start_period"] == "2025-11"
        assert observation["manual_content"].startswith("管理层补充判断")
        assert load_report_config(dataset_id)["custom_title"] == "美国汽车市场洞察周报"

        payload = build_report_payload(
            frame,
            use_llm=False,
            report_config=load_report_config(dataset_id),
            report_plan=load_report_plan(dataset_id),
        )
        assert payload["outline"][0]["subsections"][:2] == [
            "1.1 美国市场11月观察：新能源结构持续调整",
            "1.2 宏观政策动态",
        ]
        rendered = payload["market_observations"][0]
        assert [item["id"] for item in rendered["components"]] == ["market_summary", NEV_SHARE, POWER_MONTHLY]
        assert rendered["period_label"] == "2025年11月"
        assert rendered["manual_content"].startswith("管理层补充判断")
        assert all(item.get("analysis") for item in rendered["components"])

        docx_path = export_report(
            frame,
            "docx",
            tmp_path,
            "visual_composer",
            report_config=load_report_config(dataset_id),
            report_plan=load_report_plan(dataset_id),
        )
        document = Document(docx_path)
        paragraph_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        assert "1.1 美国市场11月观察：新能源结构持续调整" in paragraph_text
        assert "1.2 宏观政策动态" in paragraph_text
        assert "管理层补充判断" in paragraph_text
        assert len(document.tables) >= 1
        headers = [cell.text for cell in document.tables[0].rows[0].cells]
        assert headers[0] == "指标"
        assert "销量" in headers[1]
        assert headers[2] == "具体分析"
        assert len(document.inline_shapes) >= 2

        xlsx_path = export_report(
            frame,
            "xlsx",
            tmp_path,
            "visual_composer",
            report_config=load_report_config(dataset_id),
            report_plan=load_report_plan(dataset_id),
        )
        workbook = load_workbook(xlsx_path, read_only=True)
        summary_values = [str(cell) for row in workbook["报告目录与摘要"].iter_rows(values_only=True) for cell in row if cell]
        assert any("1.1 美国市场11月观察" in value for value in summary_values)
        assert any("管理层补充判断" in value for value in summary_values)

        pptx_path = export_report(
            frame,
            "pptx",
            tmp_path,
            "visual_composer",
            report_config=load_report_config(dataset_id),
            report_plan=load_report_plan(dataset_id),
        )
        presentation = Presentation(pptx_path)
        slide_text = "\n".join(
            shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text")
        )
        assert "1.1 美国市场11月观察：新能源结构持续调整" in slide_text
        assert "管理层补充判断" in slide_text
    finally:
        _cleanup(dataset_id)


def test_visual_composer_rejects_unavailable_component_without_writing_plan():
    dataset_id = "pytest_visual_report_invalid"
    try:
        _save(dataset_id)
        response = TestClient(app).put(
            f"/api/report-plan/{dataset_id}/visual",
            json={"observation_title": "无效图表观察", "component_ids": ["line_chart:不存在的图表"]},
        )
        assert response.status_code == 400
        assert load_report_plan(dataset_id)["market_observations"] == []
    finally:
        _cleanup(dataset_id)


def test_visual_composer_rejects_title_month_that_conflicts_with_analysis_period():
    dataset_id = "pytest_visual_report_title_period"
    try:
        _save(dataset_id)
        response = TestClient(app).put(
            f"/api/report-plan/{dataset_id}/visual",
            json={
                "observation_title": "美国市场11月观察",
                "component_ids": ["market_summary"],
                "period_mode": "single",
                "start_period": "2025-12",
                "end_period": "2025-12",
            },
        )
        assert response.status_code == 400
        assert "标题中的11月与当前分析周期2025-12不一致" in response.json()["detail"]
        assert load_report_plan(dataset_id)["market_observations"] == []
    finally:
        _cleanup(dataset_id)


def test_export_api_blocks_report_without_configured_market_observation():
    dataset_id = "pytest_visual_report_export_guard"
    try:
        _save(dataset_id)
        response = TestClient(app).post(f"/api/export/{dataset_id}/docx")
        assert response.status_code == 400
        assert "可视化编排" in response.json()["detail"]
    finally:
        _cleanup(dataset_id)
